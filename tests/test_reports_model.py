"""Importación del Modelo de Datos de Oracle Reports (XML de rwconverter)."""

from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Mm

from corporate_report_compiler.compiler import compile_project, validate_project
from corporate_report_compiler.diagnostics import Diagnostics
from corporate_report_compiler.gui import _format_result
from corporate_report_compiler.layout_reader import read_layout_template
from corporate_report_compiler.reports_convert import convert_unit, plan_conversion
from corporate_report_compiler.reports_model import load_reports_model, reports_xml
from corporate_report_compiler.workspace import format_import, import_docx, workspace_outputs


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "reports_demo.xml"
_spec = importlib.util.spec_from_file_location("layout_template_tool_r", ROOT / "tools" / "generate_layout_template.py")
tool = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = tool
_spec.loader.exec_module(tool)


def build_docx(path: Path, cells: list[str], header: str = "{{FIELD:F_NOMBRE}}", total: str = "{{SUM:F_TOTAL_PEN}}") -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    tool.run(tool.tight(document.add_paragraph(), "C"), "{{REPORT_TITLE}}")
    grid = tool.new_table(document, 1, [40, 60])
    tool.text_cell(grid.rows[0].cells[0], "Alumno :")
    tool.text_cell(grid.rows[0].cells[1], header)
    tool.spacer(document)
    tool.data_table(document, title="CUOTAS", title_span=2, headings=["N°", "Pensión", "Servicios"],
                    cells=cells, widths=[20, 30, 30], aligns=["C", "R", "R"], totals={1: total})
    raw = path.with_suffix(".raw.docx")
    document.save(raw)
    tool._sample._normalize_docx_archive(raw, path)
    raw.unlink()
    return path


def build_sections_docx(path: Path) -> Path:
    """Una tabla de Word con dos secciones de datos, como un marco de Oracle Reports:
    título subrayado, total arriba de las filas, fila de grupo con subtotal, numeración
    por fila y una segunda consulta sin filas (se oculta)."""

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    tool.run(tool.tight(document.add_paragraph(), "C"), "{{REPORT_TITLE}}")
    table = tool.new_table(document, 7, [20, 40, 30, 30])
    rows = table.rows
    title = rows[0].cells[0].merge(rows[0].cells[1])
    tool.text_cell(title, "CUOTAS PENDIENTES:", bold=True).runs[0].underline = True
    tool.text_cell(rows[0].cells[2], "Total:")
    tool.text_cell(rows[0].cells[3], "{{SUM:F_TOTAL_PEN}}", align="R")
    for cell, heading in zip(rows[1].cells, ["N°", "Ciclo", "Pensión", "Total cuota"]):
        tool.text_cell(cell, heading, bold=True)
    tool.text_cell(rows[2].cells[0], "{{FIELD:Q_CUOTAS.CICLO}}", bold=True)     # fila de grupo
    tool.text_cell(rows[2].cells[2], "{{SUM:F_TOTAL_PEN}}", align="R", bold=True)
    for cell, marker in zip(rows[3].cells, ["{{COLUMN:F_FILA}}", "{{COLUMN:Q_CUOTAS.CICLO}}", "{{COLUMN:F_PEN}}",
                                            "{{COLUMN:F_TOT_CUOTA}}"]):
        tool.text_cell(cell, marker, align="R")
    tool.text_cell(rows[4].cells[0], "Sin datos")
    tool.text_cell(rows[5].cells[0], "{{COLUMN:NRO_CUOTA2}}")
    tool.text_cell(rows[5].cells[2], "{{COLUMN:PENSION3}}")
    tool.text_cell(rows[6].cells[2], "{{SUM:PENSION3}}")
    raw = path.with_suffix(".raw.docx")
    document.save(raw)
    tool._sample._normalize_docx_archive(raw, path)
    raw.unlink()
    return path


class ReportsModelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.model = load_reports_model(FIXTURE, Diagnostics())

    def test_queries_keep_reports_names_in_select_order(self) -> None:
        self.assertEqual(self.model.queries["Q_ALUMNO"]["columns"], ["COD_ALUMNO", "NOMBRE1", "COSTO"])
        self.assertEqual(self.model.queries["Q_CUOTAS"]["columns"], ["NRO_CUOTA", "PENSION2", "SERVICIOS", "COD_ALUMNO1"])
        self.assertIn("P_COD", self.model.parameters)
        self.assertEqual(self.model.links[0]["childQuery"], "Q_CUOTAS")

    def test_field_names_and_origins_resolve_to_query_column(self) -> None:
        self.assertEqual(self.model.resolve("FIELD", "F_NOMBRE").marker, "Q_ALUMNO.NOMBRE1")
        self.assertEqual(self.model.resolve("FIELD", "nombre1").marker, "Q_ALUMNO.NOMBRE1")
        costo = self.model.resolve("FIELD", "F_COSTO")
        self.assertEqual((costo.marker, costo.mask), ("Q_ALUMNO.COSTO", "999,990.99"))   # N -> 9
        self.assertEqual(self.model.resolve("COLUMN", "F_PEN").marker, "Q_CUOTAS.PENSION2")
        self.assertEqual(self.model.resolve("FIELD", "P_COD").marker, "P_COD")

    def test_summaries_and_formulas(self) -> None:
        total = self.model.resolve("SUM", "F_TOTAL_PEN")
        self.assertEqual(total.marker, "Q_CUOTAS.PENSION2")       # total sum -> la tabla suma la columna origen
        self.assertEqual(self.model.resolve("SUM", "F_CANT").marker, "Q_CUOTAS.CS_CUOTAS")   # count: lo calcula el motor
        self.assertEqual(self.model.resolve("FIELD", "F_DOBLE").marker, "Q_ALUMNO.CF_DOBLE")
        self.assertEqual(self.model.resolve("FIELD", "F_ETIQ").marker, "Q_ALUMNO.CP_ETIQUETA")
        self.assertIn("doble(:costo)", self.model.columns["CF_DOBLE"].formula_code)
        self.assertEqual((self.model.columns["COSTO"].datatype, self.model.columns["CF_ETIQUETA"].datatype), ("N", "T"))

    def test_unknown_names(self) -> None:
        self.assertIsNone(self.model.resolve("COLUMN", "NO_EXISTE"))
        self.assertIn("F_NOMBRE", self.model.suggestions("F_NOMBR"))

    def test_xml_is_used_directly_and_other_files_are_rejected(self) -> None:
        self.assertEqual(reports_xml(FIXTURE, Diagnostics()), FIXTURE)
        diagnostics = Diagnostics()
        self.assertIsNone(reports_xml(Path("reporte.txt"), diagnostics))
        self.assertIn("REPORTS-003", {item.code for item in diagnostics.items})


class ReportsConvertTests(unittest.TestCase):
    def setUp(self) -> None:
        self.model = load_reports_model(FIXTURE, Diagnostics())

    def test_plan_converts_formulas_placeholders_summaries_and_links(self) -> None:
        plan = plan_conversion(self.model, {"NOMBRE1", "CF_DOBLE", "CP_ETIQUETA", "CF_TOTAL_CUOTA", "CS_CUOTAS"})
        self.assertEqual(plan.formulas["CF_DOBLE"], {"q": "Q_ALUMNO", "t": "N", "f": "cf_dobleformula"})
        self.assertEqual(plan.placeholders["CP_ETIQUETA"], {"t": "T", "by": ["CF_ETIQUETA"]})
        self.assertEqual(plan.summaries["CS_CUOTAS"], {"q": "Q_CUOTAS", "s": "NRO_CUOTA", "f": "count"})
        self.assertEqual(plan.filters, {"Q_CUOTAS": [{"c": "COD_ALUMNO1", "op": "eq", "p": "COD_ALUMNO"}]})
        self.assertEqual(plan.query_formulas, {"Q_ALUMNO": ["CF_DOBLE", "CF_ETIQUETA"], "Q_CUOTAS": ["CF_TOTAL_CUOTA"]})
        self.assertEqual(plan.queries, {"Q_ALUMNO", "Q_CUOTAS"})
        self.assertEqual(plan.pending, [])
        units = {unit.name: unit for unit in plan.units}
        self.assertEqual(set(units), {"cf_dobleformula", "doble", "cf_etiquetaformula", "cf_total_cuotaformula"})
        self.assertEqual(units["doble"].spec, "function doble(p_valor in number) return number;")
        self.assertIn("rr_costo NUMBER := rpt_layout.num('COSTO');", units["cf_dobleformula"].body)
        self.assertIn("return doble(rr_costo);", units["cf_dobleformula"].body)
        body = units["cf_etiquetaformula"].body
        self.assertIn("rr_cp_etiqueta VARCHAR2(4000);", body)                  # solo se asigna: sin leerlo
        self.assertIn("rr_cp_etiqueta := 'Alumno: ' || rr_nombre1; rpt_layout.set_txt('CP_ETIQUETA', rr_cp_etiqueta);", body)

    def test_unconvertible_code_becomes_a_stub(self) -> None:
        plan = plan_conversion(self.model, {"CF_AVISO"})
        unit = plan.units[0]
        self.assertIn("SRW", unit.stub)
        self.assertIn("RETURN NULL;", unit.body)
        self.assertIn("--     srw.message(100, 'aviso');", unit.body)
        self.assertIn("codigo_reports/cf_avisoformula.sql", plan.pending[0])

    def test_into_lists_strings_and_comments(self) -> None:
        code = """function CF_XFormula return Char is
  x varchar2(10);
begin
  select 'a:b', nombre into x, :CP_ETIQUETA from t where c = :P_COD; -- :NOMBRE1 no
  return x || ':COSTO';
end;"""
        result = convert_unit("cf_xformula", code, self.model)
        body = result.unit.body
        self.assertIn("into x, rr_cp_etiqueta from t where c = rr_p_cod; rpt_layout.set_txt('CP_ETIQUETA', rr_cp_etiqueta);", body)
        self.assertIn("-- :NOMBRE1 no", body)
        self.assertIn("':COSTO'", body)
        self.assertEqual((result.reads, result.assigns), ({"P_COD"}, {"CP_ETIQUETA"}))


class ReportsSectionsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.model = load_reports_model(FIXTURE, Diagnostics())

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_parent_column_binds_are_taken_from_the_parent_row(self) -> None:
        plan = plan_conversion(self.model, {"NRO_CUOTA2"})
        self.assertEqual(plan.query_refs, {"Q_VACIA": ["COD_ALUMNO"]})
        self.assertIn("Q_ALUMNO", plan.queries)                              # la consulta padre entra al proyecto

    def test_running_summary_is_detected(self) -> None:
        plan = plan_conversion(self.model, {"CS_FILA", "CS_CUOTAS"})
        self.assertEqual(plan.summaries["CS_FILA"], {"q": "Q_CUOTAS", "s": "NRO_CUOTA", "f": "count", "run": True})
        self.assertNotIn("run", plan.summaries["CS_CUOTAS"])               # total a nivel de reporte

    def test_one_word_table_with_several_queries_groups_and_totals_on_top(self) -> None:
        diagnostics = Diagnostics()
        template = read_layout_template(build_sections_docx(self.root / "s.docx"), diagnostics, self.model)
        self.assertIsNotNone(template, diagnostics.to_dict())
        self.assertNotIn("LAYOUT-STYLE-002", {item.code for item in diagnostics.items})
        blocks = [block for block in template.layout["body"] if block["k"] == "D"]
        self.assertEqual([block["q"] for block in blocks], ["Q_CUOTAS", "Q_VACIA"])
        cuotas, vacia = blocks
        self.assertEqual([row.get("g") for row in cuotas["hd"]], [None, None, ["CICLO"]])   # fila de grupo
        self.assertEqual(cuotas["sm"], ["PENSION2"])
        self.assertTrue(cuotas["hs"] and cuotas["he"])                        # total arriba; se oculta si vacía
        self.assertEqual(len(vacia["hd"]), 1)                                 # «Sin datos» es su título
        self.assertEqual(len(vacia["ft"]), 1)
        text = json.dumps(cuotas, ensure_ascii=False)
        self.assertIn('"t": "CUOTAS PENDIENTES:", "f": 2, "s": 7.0, "c": "000000", "u": 1', text)
        self.assertIn("{{COLUMN:Q_CUOTAS.CS_FILA}}", text)

    def test_group_rows_must_be_next_to_the_repeated_row(self) -> None:
        document = Document()
        document.sections[0].page_width, document.sections[0].page_height = Mm(210), Mm(297)
        table = tool.new_table(document, 3, [40, 40])
        tool.text_cell(table.rows[0].cells[0], "{{FIELD:Q_CUOTAS.NRO_CUOTA}}")
        tool.text_cell(table.rows[1].cells[0], "Encabezado")
        tool.text_cell(table.rows[2].cells[0], "{{COLUMN:F_PEN}}")
        path = self.root / "g.docx"
        document.save(self.root / "g.raw.docx")
        tool._sample._normalize_docx_archive(self.root / "g.raw.docx", path)
        diagnostics = Diagnostics()
        self.assertIsNone(read_layout_template(path, diagnostics, self.model))
        self.assertIn("LAYOUT-TABLE-007", {item.code for item in diagnostics.items})


class ReportsImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_reader_rewrites_markers_and_reports_unknown_columns(self) -> None:
        model = load_reports_model(FIXTURE, Diagnostics())
        docx = build_docx(self.root / "r.docx", ["{{COLUMN:NRO_CUOTA}}", "{{COLUMN:F_PEN}}", "{{COLUMN:NO_EXISTE}}"])
        diagnostics = Diagnostics()
        self.assertIsNone(read_layout_template(docx, diagnostics, model))
        error = next(item for item in diagnostics.items if item.code == "LAYOUT-TOKEN-009")
        self.assertIn("NO_EXISTE", error.message)

        docx = build_docx(self.root / "ok.docx", ["{{COLUMN:NRO_CUOTA}}", "{{COLUMN:F_PEN}}", "{{COLUMN:SERVICIOS|FM990D00}}"])
        diagnostics = Diagnostics()
        template = read_layout_template(docx, diagnostics, model)
        self.assertIsNotNone(template, diagnostics.to_dict())
        self.assertEqual(template.data_tables, ["Q_CUOTAS"])
        text = json.dumps(template.layout, ensure_ascii=False)
        self.assertIn("{{COLUMN:Q_CUOTAS.PENSION2|999,990.99}}", text)        # máscara tomada de Reports
        self.assertIn("{{COLUMN:Q_CUOTAS.SERVICIOS|FM990D00}}", text)         # la máscara escrita prevalece
        self.assertIn("{{SUM:Q_CUOTAS.PENSION2|999,990.00}}", text)
        self.assertIn("{{FIELD:Q_ALUMNO.NOMBRE1}}", text)

    def test_import_creates_queries_with_reports_sql_and_compiles(self) -> None:
        docx = build_docx(self.root / "boleta.docx", ["{{COLUMN:NRO_CUOTA}}", "{{COLUMN:F_PEN}}", "{{COLUMN:F_TOT_CUOTA}}"],
                          header="{{FIELD:F_ETIQ}} {{FIELD:F_DOBLE}}", total="{{SUM:F_CANT}}")
        result = import_docx(docx, projects_root=self.root / "proyectos", page="71", reports=FIXTURE)
        self.assertFalse(result.diagnostics.has_errors, result.diagnostics.to_dict())
        self.assertEqual(result.notes, ())                                     # nada pendiente a mano
        self.assertEqual(len(result.skipped), 1)                               # Q_VACIA: nadie la usa
        self.assertTrue(result.skipped[0].startswith("Q_VACIA: "), result.skipped)
        self.assertIn("no se imprime en el reporte original", result.skipped[0])
        self.assertIn("NO SE COPIARON", format_import(result, "71"))
        generated = self.root / "proyectos" / "boleta" / "generado"
        query = (generated / "q_q_cuotas.sql").read_text(encoding="utf-8")
        self.assertIn("c.pension pen", query)                                  # SQL real de Reports
        self.assertIn("COD_ALUMNO1 eq COD_ALUMNO", query)                      # Data Link: lo aplica el motor
        self.assertTrue((generated / "boleta.reports.xml").is_file())
        project = json.loads((generated / "boleta.report.json").read_text(encoding="utf-8"))
        self.assertEqual(project["reports_model"], "boleta.reports.xml")
        self.assertNotIn("reports_code", project)
        self.assertEqual(project["queries"], {"Q_ALUMNO": "q_q_alumno.sql", "Q_CUOTAS": "q_q_cuotas.sql"})
        self.assertEqual(project["parameters"], [{"name": "P_COD", "item": "P71_COD", "type": "VARCHAR2", "required": False}])

        report = generated / "boleta.report.json"
        output, process_directory = workspace_outputs(report)
        compiled = compile_project(report, output, process_directory=process_directory)
        self.assertTrue(compiled.valid, compiled.diagnostics.to_dict())
        package = (self.root / "proyectos" / "boleta" / "rpt_boleta.sql").read_text(encoding="utf-8")
        self.assertIn("l_names.append('PENSION2');", package)
        self.assertIn("l_query.put('columns', l_names);", package)
        spec = package[:package.index("CREATE OR REPLACE PACKAGE BODY")]
        self.assertIn("    function CF_DOBLEFormula return Number;", spec)
        self.assertIn("    function doble(p_valor in number) return number;", spec)
        self.assertIn("rr_costo NUMBER := rpt_layout.num('COSTO');", package)
        self.assertIn("""l_query.put('filters', JSON_ARRAY_T.parse(q'~[{"c":"COD_ALUMNO1","op":"eq","p":"COD_ALUMNO"}]~'));""", package)
        self.assertIn("""l_query.put('formulas', JSON_ARRAY_T.parse(q'~["CF_TOTAL_CUOTA"]~'));""", package)
        self.assertIn("p_model   => ", package)
        layout = json.loads((generated / "layout.json").read_text(encoding="utf-8"))
        self.assertEqual(layout["reports"]["model"]["package"], "RPT_BOLETA")
        self.assertIn("{{SUM:Q_CUOTAS.CS_CUOTAS}}", json.dumps(layout["layout"]))

    def test_unconvertible_formula_gets_a_file_to_fix(self) -> None:
        docx = build_docx(self.root / "f.docx", ["{{COLUMN:NRO_CUOTA}}", "{{COLUMN:F_PEN}}", "{{COLUMN:SERVICIOS}}"],
                          header="{{FIELD:CF_AVISO}}")
        result = import_docx(docx, projects_root=self.root / "proyectos", page="71", reports=FIXTURE)
        self.assertTrue(any("codigo_reports/cf_avisoformula.sql" in note for note in result.notes), result.notes)
        generated = self.root / "proyectos" / "f" / "generado"
        fix = generated / "codigo_reports" / "cf_avisoformula.sql"
        self.assertIn("srw.message(100, 'aviso');", fix.read_text(encoding="utf-8"))
        project = generated / "f.report.json"
        self.assertEqual(json.loads(project.read_text(encoding="utf-8"))["reports_code"], "codigo_reports")
        validated = validate_project(project)
        self.assertTrue(validated.valid, validated.diagnostics.to_dict())
        self.assertIn("LAYOUT-REPORTS-011", {item.code for item in validated.diagnostics.items})
        summary = _format_result(validated)                                    # lo que muestra la GUI
        self.assertTrue(summary.startswith("PÁGINA APEX Y PAGE ITEMS\n" + "=" * 24 + "\nPágina elegida: 71\n"))
        self.assertIn("  P71_COD  VARCHAR2  -> :P_COD", summary)                     # recordatorio de Page Items
        self.assertIn("CÓDIGO DE ORACLE REPORTS", summary)
        self.assertIn("FALTA INSERTAR O REVISAR:", summary)
        self.assertIn("codigo_reports/cf_avisoformula.sql", summary)

        # El usuario corrige la función con la sintaxis de Reports: deja de estar pendiente.
        fix.write_text(fix.read_text(encoding="utf-8").replace("srw.message(100, 'aviso');", "null;"),
                       encoding="utf-8")
        validated = validate_project(project)
        self.assertNotIn("LAYOUT-REPORTS-011", {item.code for item in validated.diagnostics.items})
        self.assertIn("Nada pendiente a mano.", _format_result(validated))
        unit = next(u for u in validated.definition["reports"]["units"] if u["name"] == "cf_avisoformula")
        self.assertIn("null;", unit["body"])
        # Volver a cargar el Word conserva la corrección.
        import_docx(docx, projects_root=self.root / "proyectos", page="71", reports=FIXTURE, replace=True)
        self.assertIn("null;", fix.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()

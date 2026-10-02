"""Los entregables Word contienen las instrucciones críticas del contrato."""

from __future__ import annotations

import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REQUIRED = {
    "deliverables/Manual_de_uso.docx": (
        "Reload on Submit",
        "Always",
        "Submit Page",
        "When Button Pressed",
        "P0_REPORT_FORMAT",
        "Compilador de reportes APEX",
        "Nuevo proyecto desde DOCX",
        "SQL-013",
        "automáticamente",
        "{{FIELD:UNIDAD}}",
    ),
    "deliverables/Informe_de_alcance_y_limitaciones.docx": ("chr(38)", "bytes UTF-8"),
    "docs/CONTRACT.docx": ("bytes UTF-8", "ecuaciones", "nombres reservados", "chr(38)", "automáticamente"),
    "docs/RELEASE_REVIEW.docx": ("GO para empaquetado", "Reload on Submit"),
    "docs/TEMPLATE_QA.docx": ("templates", "{{FIELD:UNIDAD}}"),
    "README.docx": ("Abrir compilador.bat", "automáticamente"),
    "sql/README.docx": ("Reload on Submit", "-20173"),
    "examples/README.docx": ("build\\entidades","{{FIELD:UNIDAD}}"),
}
# Nombres de entregas previas: la documentación describe solo la versión actual.
FORBIDDEN = (
    "DOWNLOAD_IG",
    "v7",
    "versión 8",
    "p_display_columns_json",
    "p_column_spans_json",
    "p_template_json",
    "exclude_columns",
    "alias obsoleto",
    "{{VALUE:",
)
MARKERS = ("{{REPORT_TITLE}}", "{{FIELD:", "{{COLUMN:", "{{APP_USER}}", "{{GENERATED_AT}}")
TEMPLATES = (
    "templates/entidades_portrait.docx",
    "templates/entidades_landscape.docx",
    "examples/entidades.docx",
)


def document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


class DeliverableDocumentationTests(unittest.TestCase):
    def test_word_documents_contain_critical_instructions(self) -> None:
        for relative, phrases in REQUIRED.items():
            with self.subTest(document=relative):
                path = ROOT / relative
                self.assertTrue(path.is_file(), f"Falta {relative}")
                text = document_text(path)
                for phrase in phrases:
                    self.assertIn(phrase, text, f"{relative} no contiene {phrase!r}")

    def test_word_documents_omit_previous_versions(self) -> None:
        for relative in REQUIRED:
            text = document_text(ROOT / relative)
            for name in FORBIDDEN:
                with self.subTest(document=relative, name=name):
                    self.assertNotIn(name, text)

    def test_word_documents_include_flow_diagrams(self) -> None:
        for relative in REQUIRED:
            with self.subTest(document=relative):
                self.assertIn("Figura 1.", document_text(ROOT / relative))

    def test_reference_templates_use_every_marker_kind(self) -> None:
        for relative in TEMPLATES:
            text = document_text(ROOT / relative)
            for marker in MARKERS:
                with self.subTest(template=relative, marker=marker):
                    self.assertIn(marker, text)


if __name__ == "__main__":
    unittest.main()

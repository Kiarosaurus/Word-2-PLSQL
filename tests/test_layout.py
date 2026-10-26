"""Modo layout: varias tablas, maquetación, totales y número de página."""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Mm

from corporate_report_compiler.compiler import compile_project, validate_project
from corporate_report_compiler.diagnostics import Diagnostics
from corporate_report_compiler.layout_reader import is_layout_docx, read_layout_template
from corporate_report_compiler.workspace import workspace_outputs


ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "proyectos" / "demo" / "estado_cuenta"
_spec = importlib.util.spec_from_file_location("layout_template_tool", ROOT / "tools" / "generate_layout_template.py")
tool = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = tool
_spec.loader.exec_module(tool)


def save(document: Document, path: Path) -> Path:
    raw = path.with_suffix(".raw.docx")
    document.save(raw)
    tool._sample._normalize_docx_archive(raw, path)
    raw.unlink()
    return path


def base_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    return document


def data_table(document: Document, body: list[str], footer: list[str] | None = None) -> None:
    table = tool.new_table(document, 3 if footer else 2, [30] * len(body))
    for index, marker in enumerate(body):
        tool.text_cell(table.rows[0].cells[index], f"Col {index}")
        tool.text_cell(table.rows[1].cells[index], marker)
        if footer:
            tool.text_cell(table.rows[2].cells[index], footer[index])


def codes(diagnostics: Diagnostics) -> set[str]:
    return {item.code for item in diagnostics.items}


class LayoutReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def read(self, document: Document) -> tuple[object, Diagnostics]:
        diagnostics = Diagnostics(strict=True)
        path = save(document, self.root / "t.docx")
        return read_layout_template(path, diagnostics), diagnostics

    def test_demo_template_structure(self) -> None:
        diagnostics = Diagnostics(strict=True)
        template = read_layout_template(DEMO / "estado_cuenta.docx", diagnostics)
        self.assertIsNotNone(template, diagnostics.to_dict())
        self.assertTrue(is_layout_docx(DEMO / "estado_cuenta.docx"))
        self.assertEqual(template.data_tables, ["CAJA", "CUOTAS", "TRANSFERENCIAS", "ATRASOS"])
        self.assertEqual({"SISTEMA", "MODULO", "CODIGO"}, template.fields)
        layout = template.layout
        self.assertEqual([block["k"] for block in layout["header"]], ["G", "P"])
        self.assertEqual([block["k"] for block in layout["footer"]], ["G"])
        cuotas = next(block for block in layout["body"] if block.get("q") == "CUOTAS")
        self.assertEqual(cuotas["sm"], ["PENSION", "SERVICIOS", "TOTAL"])
        self.assertEqual(len(cuotas["hd"]), 2)          # banda + títulos, se repiten por página
        grid = layout["body"][0]
        self.assertEqual(grid["k"], "G")
        band = grid["rs"][0]["cs"][0]
        self.assertEqual((band["sp"], band["fl"]), (8, "D9D9D9"))
        # Bordes blancos (invisibles) y lados sin borde coexisten en la cuadrícula.
        label = grid["rs"][1]["cs"][0]
        self.assertEqual(label["b"]["r"][1], "FFFFFF")
        self.assertNotIn("t", label["b"])

    def test_reports_hierarchy_uses_one_detail_pass_and_group_subtotals(self) -> None:
        project = ROOT / "proyectos" / "FCR118_E" / "generado" / "FCR118_E.report.json"
        result = validate_project(project)
        self.assertTrue(result.valid, result.diagnostics.to_dict())
        tables = [block for block in result.definition["layout"]["body"] if block["k"] == "D"]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual([block["k"] for block in result.definition["layout"]["header"]], ["G", "G"])
        self.assertEqual(len(result.definition["layout"]["header"][1]["rs"]), 2)
        self.assertEqual(table["vis"], [
            ["MOV_CAR", "MOV_ABO", "CF_ACU_CAR_ANT", "CF_ACU_ABO_ANT"],
            ["CF_SALDO_INI8", "CF_SALDO_FIN8"],
        ])
        self.assertTrue(table["vg"])
        self.assertEqual(
            [row["g"] for row in table["hd"]],
            [
                ["NORDEN", "CTA2"],
                ["NORDEN", "CTA2", "CTA3"],
                ["NORDEN", "CTA2", "CTA3", "CTA4"],
            ],
        )
        for row in table["hd"]:
            subtotal = row["cs"][1]["ps"][0]["ln"][0][0]["t"]
            self.assertEqual(subtotal, "{{SUM:Q_1.CF_ACU_CAR_ANT|999,999,999.99}}")
        self.assertEqual(table["bd"]["cs"][0]["sp"], 2)
        self.assertEqual(table["bd"]["cs"][0]["ps"][0]["ln"][0][0]["t"], "{{COLUMN:Q_1.CTA8}}")
        self.assertIn("{{FIELD:Q_1.CF_UTI_PAS", str(result.definition["layout"]["body"]))
        self.assertNotIn("CS_UTI_PAS", result.definition["constants"])
        self.assertEqual(
            {p["name"]: p["default"] for p in result.definition["parameters"] if "default" in p},
            {"NUMERO1": 303, "NUMERO2": 306, "NUMERO3": 63},
        )
        self.assertIn("CF_SALDO_INI8", result.definition["reports"]["model"]["formulas"])
        self.assertIn("CF_SALDO_FIN8", result.definition["reports"]["model"]["formulas"])

    def test_simple_templates_are_not_layout(self) -> None:
        self.assertFalse(is_layout_docx(ROOT / "templates" / "entidades_portrait.docx"))

    def test_page_number_only_in_header_or_footer(self) -> None:
        document = base_document()
        tool.run(document.add_paragraph(), "Página {{PAGE}}")
        data_table(document, ["{{COLUMN:Q.A}}"])
        _template, diagnostics = self.read(document)
        self.assertIn("LAYOUT-TOKEN-003", codes(diagnostics))

    def test_column_marker_rules(self) -> None:
        document = base_document()
        tool.run(document.add_paragraph(), "{{COLUMN:Q.A}}")
        data_table(document, ["{{COLUMN:Q.A}}", "{{COLUMN:OTRA.B}}"], ["{{SUM:Q.A}}", "{{SUM:OTRA.B}}"])
        _template, diagnostics = self.read(document)
        self.assertIn("LAYOUT-TOKEN-005", codes(diagnostics))      # COLUMN fuera de una tabla de datos
        self.assertIn("LAYOUT-TOKEN-006", codes(diagnostics))      # dos consultas en la misma tabla
        misplaced = next(item for item in diagnostics.items if item.code == "LAYOUT-TOKEN-005")
        self.assertIn("Se encontró", misplaced.message)
        self.assertIn("se esperaba", misplaced.message)
        self.assertIsNotNone(misplaced.suggestion)

    def test_vertical_merge_rejected(self) -> None:
        document = base_document()
        table = tool.new_table(document, 2, [30, 30])
        table.cell(0, 0).merge(table.cell(1, 0))
        data_table(document, ["{{COLUMN:Q.A}}"])
        _template, diagnostics = self.read(document)
        self.assertIn("LAYOUT-TABLE-004", codes(diagnostics))

    def test_rows_shorter_than_the_table_are_accepted(self) -> None:
        # Word: fila acortada arrastrando su borde -> columnas vacías gridBefore/gridAfter.
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = base_document()
        table = tool.new_table(document, 2, [20, 20, 20])
        rows = table._tbl.findall(qn("w:tr"))
        for row, (name, keep) in zip(rows, (("w:gridAfter", 0), ("w:gridBefore", 2))):
            cells = row.findall(qn("w:tc"))
            for index, cell in enumerate(cells):
                if index != keep:
                    row.remove(cell)
            props = row.find(qn("w:trPr"))
            if props is None:
                props = OxmlElement("w:trPr")
                row.insert(0, props)
            gap = OxmlElement(name)
            gap.set(qn("w:val"), "2")
            props.append(gap)
        data_table(document, ["{{COLUMN:Q.A}}"])
        template, diagnostics = self.read(document)
        self.assertNotIn("LAYOUT-TABLE-005", codes(diagnostics))
        self.assertNotIn("LAYOUT-TABLE-003", codes(diagnostics))
        grid = next(block for block in template.layout["body"] if block.get("k") == "G")
        self.assertEqual([cell["c"] for cell in grid["rs"][0]["cs"]], [0])        # hueco al final
        self.assertEqual([cell["c"] for cell in grid["rs"][1]["cs"]], [2])        # hueco al inicio

    def test_floating_table_is_warned(self) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = base_document()
        table = tool.new_table(document, 1, [30, 30])
        position = OxmlElement("w:tblpPr")
        position.set(qn("w:vertAnchor"), "text")
        position.set(qn("w:tblpY"), "-4283")
        table._tbl.tblPr.append(position)
        data_table(document, ["{{COLUMN:Q.A}}"])
        _template, diagnostics = self.read(document)
        warning = next(item for item in diagnostics.items if item.code == "LAYOUT-TABLE-006")
        self.assertEqual(warning.severity, "WARNING")
        self.assertIn("Ajuste del texto: Ninguno", warning.suggestion)

    def test_empty_markers_from_ai_are_listed_as_pending(self) -> None:
        document = base_document()
        tool.new_table(document, 1, [30, 30])
        tool.text_cell(document.tables[0].rows[0].cells[0], "Código : {{FIELD:}}")
        data_table(document, ["{{COLUMN:}}", "{{COLUMN:}}"], ["", "{{SUM:}}"])
        path = save(document, self.root / "ia.docx")

        self.assertTrue(is_layout_docx(path))        # varias tablas: modo layout aunque falten nombres
        diagnostics = Diagnostics(strict=True)
        self.assertIsNone(read_layout_template(path, diagnostics))
        pending = [item for item in diagnostics.items if item.code == "LAYOUT-TOKEN-008"]
        self.assertEqual(len(pending), 4)
        self.assertEqual(codes(diagnostics), {"LAYOUT-TOKEN-008"})
        self.assertIn("CONSULTA.COLUMNA", pending[0].suggestion)


class LayoutProjectTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.copy = Path(self.temporary.name) / "estado_cuenta"
        shutil.copytree(DEMO, self.copy)
        self.project = self.copy / "generado" / "estado_cuenta.report.json"

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def edit(self, change) -> None:
        raw = json.loads(self.project.read_text(encoding="utf-8"))
        change(raw)
        self.project.write_text(json.dumps(raw), encoding="utf-8")

    def test_demo_project_validates(self) -> None:
        result = validate_project(self.project)
        self.assertTrue(result.valid, result.diagnostics.to_dict())
        self.assertEqual(result.kind, "layout")

    def test_undeclared_bind_missing_query_and_unknown_field(self) -> None:
        def change(raw: dict) -> None:
            raw["parameters"] = []
            del raw["queries"]["ATRASOS"]
            del raw["constants"]["SISTEMA"]
        self.edit(change)
        found = codes(validate_project(self.project).diagnostics)
        self.assertIn("LAYOUT-PROJECT-022", found)
        self.assertIn("LAYOUT-PROJECT-040", found)
        self.assertIn("LAYOUT-PROJECT-042", found)

    def test_compile_places_upload_files_next_to_docx(self) -> None:
        output, process_directory = workspace_outputs(self.project)
        result = compile_project(self.project, output, process_directory=process_directory)
        self.assertTrue(result.valid, result.diagnostics.to_dict())
        names = {path.relative_to(self.copy).as_posix() for path in result.artifacts}
        self.assertEqual(names, {"apex_process.sql", "rpt_estado_cuenta.sql", "generado/layout.json", "generado/validation.json"})
        package = (self.copy / "rpt_estado_cuenta.sql").read_text(encoding="utf-8")
        self.assertIn("FUNCTION build(", package)
        self.assertIn("p_cod_alumno IN VARCHAR2", package)
        self.assertIn("FUNCTION q_cuotas RETURN CLOB", package)
        self.assertIn("rpt_layout.render(", package)
        self.assertIn(":P71_COD_ALUMNO", (self.copy / "apex_process.sql").read_text(encoding="utf-8"))

    def test_committed_demo_outputs_match_a_fresh_compilation(self) -> None:
        output, process_directory = workspace_outputs(self.project)
        compile_project(self.project, output, process_directory=process_directory)
        for relative in ("apex_process.sql", "rpt_estado_cuenta.sql", "generado/layout.json"):
            self.assertEqual((self.copy / relative).read_bytes(), (DEMO / relative).read_bytes(), relative)

    def test_pdf_is_downloaded_by_default_and_can_open_inline(self) -> None:
        from corporate_report_compiler.apex_guide import process_code

        self.assertIn("Content-Disposition: attachment; filename=\"estado_cuenta.pdf\"",
                      process_code(validate_project(self.project)))
        raw = json.loads(self.project.read_text(encoding="utf-8"))
        raw["disposition"] = "inline"
        self.project.write_text(json.dumps(raw), encoding="utf-8")
        self.assertIn("Content-Disposition: inline;", process_code(validate_project(self.project)))
        raw["disposition"] = "abrir"
        self.project.write_text(json.dumps(raw), encoding="utf-8")
        self.assertIn("LAYOUT-PROJECT-007", {item.code for item in validate_project(self.project).diagnostics.items})


if __name__ == "__main__":
    unittest.main()

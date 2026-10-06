from __future__ import annotations

import json
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

from corporate_report_compiler.compiler import compile_project, validate_project

from tests.support import make_valid_project


ROOT = Path(__file__).resolve().parents[1]
PACKAGE_PATH = ROOT / "sql" / "modo_simple" / "pkg_corporate_reports.sql"
EXPECTED_ARTIFACTS = ("apex_process.sql", "template.json", "validation.json")


def _download_query_signature() -> tuple[set[str], set[str]]:
    """Return all and required DOWNLOAD_QUERY parameter names from the spec."""

    source = PACKAGE_PATH.read_text(encoding="utf-8")
    spec_match = re.search(
        r"CREATE\s+OR\s+REPLACE\s+PACKAGE\s+pkg_corporate_reports"
        r"(?:\s+AUTHID\s+(?:CURRENT_USER|DEFINER))?\s+AS"
        r"(?P<spec>.*?)END\s+pkg_corporate_reports\s*;\s*/",
        source,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if spec_match is None:
        raise AssertionError("No se encontró la especificación del package.")
    procedure_match = re.search(
        r"PROCEDURE\s+download_query\s*\((?P<args>.*?)\)\s*;",
        spec_match.group("spec"),
        flags=re.IGNORECASE | re.DOTALL,
    )
    if procedure_match is None:
        raise AssertionError("No se encontró DOWNLOAD_QUERY en la especificación.")

    all_names: set[str] = set()
    required: set[str] = set()
    for declaration in procedure_match.group("args").split(","):
        name_match = re.match(r"\s*(p_[a-z0-9_]+)\s+in\b", declaration, re.IGNORECASE)
        if name_match is None:
            continue
        name = name_match.group(1).lower()
        all_names.add(name)
        if not re.search(r"\bdefault\b", declaration, re.IGNORECASE):
            required.add(name)
    return all_names, required


def _generated_call_parameters(process: str) -> set[str]:
    call = re.search(
        r"pkg_corporate_reports\.download_query\s*\((?P<args>.*?)\)\s*;",
        process,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if call is None:
        raise AssertionError("El proceso generado no llama a DOWNLOAD_QUERY.")
    return {
        name.lower()
        for name in re.findall(r"\b(p_[a-z0-9_]+)\s*=>", call.group("args"), re.IGNORECASE)
    }


class CompilerIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_complete_project_compiles_to_verified_apex_artifacts(self) -> None:
        project_path = make_valid_project(self.root / "project")
        output = self.root / "compiled"

        result = compile_project(project_path, output)

        self.assertTrue(result.success, result.diagnostics.to_dict())
        self.assertEqual(
            tuple(sorted(path.name for path in result.output_files)),
            EXPECTED_ARTIFACTS,
        )
        self.assertEqual(tuple(sorted(path.name for path in output.iterdir())), EXPECTED_ARTIFACTS)

        definition = json.loads((output / "template.json").read_text(encoding="utf-8"))
        validation = json.loads((output / "validation.json").read_text(encoding="utf-8"))
        process = (output / "apex_process.sql").read_text(encoding="utf-8")

        self.assertEqual(definition["schema_version"], "1.0")
        self.assertEqual(definition["report"]["id"], "PERSONAS")
        self.assertEqual(
            [column["name"] for column in definition["columns"]],
            ["VDNI", "VNOM", "DEPARTAMENTO"],
        )
        self.assertEqual(definition["query"]["bindings"][0]["bind"], "P_DNI")
        self.assertEqual(definition["fields"][0]["name"], "OFFICE_NAME")
        self.assertTrue(validation["valid"])
        self.assertEqual(validation["errors"], 0)

        package_parameters, required_parameters = _download_query_signature()
        generated_parameters = _generated_call_parameters(process)
        forbidden_parameters = {"p_paper_size", "p_auto_reserve_pct"}
        self.assertTrue(
            forbidden_parameters.isdisjoint(package_parameters),
            "DOWNLOAD_QUERY expone parámetros internos del layout: "
            f"{sorted(forbidden_parameters & package_parameters)}",
        )
        self.assertTrue(
            forbidden_parameters.isdisjoint(generated_parameters),
            "El proceso generado intenta configurar decisiones internas: "
            f"{sorted(forbidden_parameters & generated_parameters)}",
        )
        self.assertFalse(
            generated_parameters - package_parameters,
            "El proceso usa parámetros que no existen en DOWNLOAD_QUERY: "
            f"{sorted(generated_parameters - package_parameters)}",
        )
        self.assertFalse(
            required_parameters - generated_parameters,
            "El proceso omite parámetros obligatorios de DOWNLOAD_QUERY: "
            f"{sorted(required_parameters - generated_parameters)}",
        )

    def test_compilation_is_byte_for_byte_deterministic(self) -> None:
        project_path = make_valid_project(self.root / "project")
        first = self.root / "first"
        second = self.root / "second"

        first_result = compile_project(project_path, first)
        second_result = compile_project(project_path, second)

        self.assertTrue(first_result.success, first_result.diagnostics.to_dict())
        self.assertTrue(second_result.success, second_result.diagnostics.to_dict())
        for name in EXPECTED_ARTIFACTS:
            with self.subTest(artifact=name):
                self.assertEqual((first / name).read_bytes(), (second / name).read_bytes())

    def test_unspecified_word_alignment_reaches_runtime_as_null(self) -> None:
        project_path = make_valid_project(self.root / "automatic-alignment")
        payload = json.loads(project_path.read_text(encoding="utf-8"))
        template_path = project_path.parent / payload["template"]
        document = Document(template_path)
        document.tables[0].cell(1, 0).paragraphs[0].alignment = None
        document.save(template_path)

        output = self.root / "automatic-alignment-output"
        result = compile_project(project_path, output)

        self.assertTrue(result.success, result.diagnostics.to_dict())
        definition = json.loads((output / "template.json").read_text(encoding="utf-8"))
        self.assertIsNone(definition["columns"][0]["alignment"])
        process = (output / "apex_process.sql").read_text(encoding="utf-8")
        self.assertIn('"alignment":null', process)

    def test_unspecified_column_widths_use_contractual_equal_weights(self) -> None:
        project_path = make_valid_project(
            self.root / "equal-default-widths",
            project_overrides={"column_widths": []},
        )
        output = self.root / "equal-default-widths-output"
        result = compile_project(project_path, output)

        self.assertTrue(result.success, result.diagnostics.to_dict())
        definition = json.loads((output / "template.json").read_text(encoding="utf-8"))
        self.assertEqual(
            [column["width"] for column in definition["columns"]],
            [
                {"mode": "WEIGHT", "value": 1.0},
                {"mode": "WEIGHT", "value": 1.0},
                {"mode": "WEIGHT", "value": 1.0},
            ],
        )

    def test_validation_never_writes_and_invalid_compile_leaves_no_partial_output(self) -> None:
        project_path = make_valid_project(
            self.root / "project",
            project_overrides={
                "excluded_columns": ["VDNI", "VNOM", "DEPARTAMENTO"],
            },
        )
        output = self.root / "must-not-exist"

        validation = validate_project(project_path)
        compilation = compile_project(project_path, output)

        self.assertFalse(validation.valid)
        self.assertFalse(compilation.success)
        self.assertEqual(validation.artifacts, ())
        self.assertEqual(compilation.artifacts, ())
        self.assertFalse(output.exists(), "Una compilación inválida no debe crear la salida.")

    def test_invalid_recompile_preserves_last_known_good_output(self) -> None:
        project_path = make_valid_project(self.root / "project")
        output = self.root / "compiled"
        first = compile_project(project_path, output)
        self.assertTrue(first.success, first.diagnostics.to_dict())
        before = {name: (output / name).read_bytes() for name in EXPECTED_ARTIFACTS}

        payload = json.loads(project_path.read_text(encoding="utf-8"))
        payload["excluded_columns"] = ["VDNI", "VNOM", "DEPARTAMENTO"]
        project_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        second = compile_project(project_path, output)

        self.assertFalse(second.success)
        self.assertEqual(second.artifacts, ())
        self.assertEqual(
            {name: (output / name).read_bytes() for name in EXPECTED_ARTIFACTS},
            before,
        )


if __name__ == "__main__":
    unittest.main()

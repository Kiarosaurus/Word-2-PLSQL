"""Reproducible test fixtures for DOCX and compiler projects.

The helpers intentionally generate only synthetic data.  Security fixtures are
created by mutating the OOXML/ZIP package after python-docx has saved a valid
document, which makes each hostile property explicit in its test.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DEFAULT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("DNI", "VDNI"),
    ("Nombre completo", "VNOM"),
    ("Departamento", "DEPARTAMENTO"),
)

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _remove_implicit_custom_xml(source: Path, destination: Path) -> None:
    """Remove the bibliography customXml bundled by python-docx.

    The production scanner rejects every customXml part. Fixtures therefore
    start from the actual public contract; hostile customXml is added only by
    the security test that intends to exercise that rejection.
    """

    def sanitize(name: str, data: bytes) -> bytes:
        if name.casefold().endswith(".rels"):
            root = ET.fromstring(data)
            for relationship in list(root):
                relation_type = relationship.attrib.get("Type", "").casefold()
                target = relationship.attrib.get("Target", "").replace("\\", "/").casefold()
                if relation_type.endswith("/customxml") or "customxml/" in target:
                    root.remove(relationship)
            ET.register_namespace("", REL_NS)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        if name == "[Content_Types].xml":
            root = ET.fromstring(data)
            for child in list(root):
                if child.attrib.get("PartName", "").casefold().startswith("/customxml/"):
                    root.remove(child)
            ET.register_namespace("", CT_NS)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return data

    with zipfile.ZipFile(source, "r") as input_zip:
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as output_zip:
            for info in input_zip.infolist():
                if info.filename.casefold().startswith("customxml/"):
                    continue
                output_zip.writestr(info, sanitize(info.filename, input_zip.read(info.filename)))


def _set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001 - OOXML fixture helper
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def _set_cell_border(cell, color: str = "BFC3C7", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # noqa: SLF001 - OOXML fixture helper
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _style_run(
    run,
    *,
    font: str = "Arial",
    size: float = 9,
    bold: bool = False,
    color: str = "25282B",
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_valid_docx(
    path: Path,
    *,
    columns: Iterable[tuple[str, str]] = DEFAULT_COLUMNS,
    orientation: str = "PORTRAIT",
    include_field: bool = True,
    include_builtins: bool = True,
    font: str = "Arial",
    split_placeholder_runs: bool = False,
) -> Path:
    """Build a minimal template conforming to contract 1.0."""

    columns = tuple(columns)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    if orientation.upper() == "LANDSCAPE":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = header.add_run("{{REPORT_TITLE}}")
    _style_run(title, font=font, size=15, bold=True, color="2F343A")
    if include_field:
        extra = header.add_run("\nOficina: {{FIELD:OFFICE_NAME}}")
        _style_run(extra, font=font, size=15, bold=True, color="2F343A")

    table = document.add_table(rows=2, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(8.0), Cm(5.0)]
    for index, (heading, alias) in enumerate(columns):
        head_cell = table.cell(0, index)
        body_cell = table.cell(1, index)
        for cell in (head_cell, body_cell):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_border(cell)
            cell.width = widths[index] if index < len(widths) else Cm(3.0)

        head_cell.text = ""
        head_paragraph = head_cell.paragraphs[0]
        head_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        head_run = head_paragraph.add_run(heading)
        _style_run(head_run, font=font, size=9, bold=True, color="FFFFFF")
        _set_cell_fill(head_cell, "4A4F55")

        body_cell.text = ""
        body_paragraph = body_cell.paragraphs[0]
        body_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        marker = f"{{{{COLUMN:{alias}}}}}"
        if split_placeholder_runs:
            split_at = max(2, len(marker) // 2)
            for fragment in (marker[:split_at], marker[split_at:]):
                run = body_paragraph.add_run(fragment)
                _style_run(run, font=font, size=8.5, color="25282B")
        else:
            body_run = body_paragraph.add_run(marker)
            _style_run(body_run, font=font, size=8.5, color="25282B")
        _set_cell_fill(body_cell, "FFFFFF")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = (
        "Generado por {{APP_USER}} el {{GENERATED_AT}}"
        if include_builtins
        else "Reporte corporativo"
    )
    footer_run = footer.add_run(footer_text)
    _style_run(footer_run, font=font, size=8, color="666666")

    path.parent.mkdir(parents=True, exist_ok=True)
    raw_path = path.with_name(f".{path.name}.raw")
    try:
        document.save(raw_path)
        _remove_implicit_custom_xml(raw_path, path)
    finally:
        raw_path.unlink(missing_ok=True)
    return path


def write_project(
    directory: Path,
    *,
    template_name: str = "template.docx",
    query: str | None = None,
    columns: Iterable[tuple[str, str]] = DEFAULT_COLUMNS,
    project_overrides: dict | None = None,
    include_field: bool = True,
) -> Path:
    """Write a valid project JSON and SQL next to a template."""

    directory.mkdir(parents=True, exist_ok=True)
    columns = tuple(columns)
    query = query if query is not None else (
        "select p.vdni, p.vnom, d.descripcion as departamento\n"
        "  from persona p\n"
        "  left join departamento d on d.id = p.departamento_id\n"
        " where (:P_DNI is null or p.vdni = :P_DNI)\n"
        " order by p.vnom"
    )
    query_path = directory / "report.sql"
    query_path.write_text(query + "\n", encoding="utf-8")

    project = {
        "schema": "corporate-report-project/1.0",
        "report_id": "PERSONAS",
        "template": template_name,
        "query_file": query_path.name,
        "title": "Relación de personas",
        "file_name": "personas",
        "max_rows": 2000,
        "bindings": [
            {
                "name": "P_DNI",
                "item": "P42_DNI",
                "type": "VARCHAR2",
                "required": False,
            }
        ],
        "fields": (
            [
                {
                    "name": "OFFICE_NAME",
                    "source": "ITEM",
                    "item": "P42_OFFICE_NAME",
                    "type": "VARCHAR2",
                }
            ]
            if include_field
            else []
        ),
        "excluded_columns": [],
        "column_widths": [
            (
                {"column": alias, "mode": "FIXED_PERCENT", "value": 14}
                if index == 0 and len(columns) > 1
                else {"column": alias, "mode": "AUTO"}
                if index == min(1, len(columns) - 1)
                else {"column": alias, "mode": "WEIGHT", "value": 1}
            )
            for index, (_heading, alias) in enumerate(columns)
        ],
    }
    if project_overrides:
        project.update(project_overrides)

    project_path = directory / "personas.report.json"
    project_path.write_text(
        json.dumps(project, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return project_path


def make_valid_project(
    directory: Path,
    *,
    columns: Iterable[tuple[str, str]] = DEFAULT_COLUMNS,
    query: str | None = None,
    project_overrides: dict | None = None,
    include_field: bool = True,
    orientation: str = "PORTRAIT",
    font: str = "Arial",
    split_placeholder_runs: bool = False,
) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    template_path = directory / "template.docx"
    build_valid_docx(
        template_path,
        columns=columns,
        orientation=orientation,
        include_field=include_field,
        font=font,
        split_placeholder_runs=split_placeholder_runs,
    )
    return write_project(
        directory,
        query=query,
        columns=columns,
        project_overrides=project_overrides,
        include_field=include_field,
    )


def rewrite_zip(
    source: Path,
    destination: Path,
    *,
    replacements: dict[str, bytes] | None = None,
    additions: dict[str, bytes] | None = None,
    omit: set[str] | None = None,
) -> Path:
    """Copy a DOCX ZIP while replacing, adding or omitting named entries."""

    replacements = replacements or {}
    additions = additions or {}
    omit = omit or set()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "r") as source_zip:
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as target_zip:
            for info in source_zip.infolist():
                if info.filename in omit:
                    continue
                data = replacements.get(info.filename, source_zip.read(info.filename))
                target_zip.writestr(info, data)
            for name, data in additions.items():
                target_zip.writestr(name, data)
    return destination


def replace_template_in_project(project_path: Path, template_path: Path) -> None:
    project = json.loads(project_path.read_text(encoding="utf-8"))
    project["template"] = template_path.name
    project_path.write_text(
        json.dumps(project, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def copy_project_tree(source: Path, destination: Path) -> Path:
    shutil.copytree(source, destination)
    return destination / "personas.report.json"

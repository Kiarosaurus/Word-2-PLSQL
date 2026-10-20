"""Tablas de los FROM de un proyecto y su esquema / DB link."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from corporate_report_compiler.compiler import validate_project
from corporate_report_compiler.sources import common_source, project_tables, qualified, set_table_source, table_spans
from corporate_report_compiler.workspace import import_docx

from tests.test_reports_model import FIXTURE, build_docx


class TableSpansTests(unittest.TestCase):
    def test_finds_tables_of_from_and_join_only(self) -> None:
        sql = """select a.x, extract(year from a.f), trim(leading '0' from b.c),
                        (select max(z) from zz) m, 'from texto' t
                   from pgpago a, bdatos.presol@DL_11G b
                   left join cfperofi c on c.id = a.id
                   join table(f(1)) t on 1 = 1
                  where exists (select 1 from dual)  -- from comentario
                 union all select * from Otra.TABLA o"""
        self.assertEqual([(span.table, span.text) for span in table_spans(sql)], [
            ("ZZ", "zz"), ("PGPAGO", "pgpago"), ("PRESOL", "bdatos.presol@DL_11G"),
            ("CFPEROFI", "cfperofi"), ("TABLA", "Otra.TABLA"),
        ])

    def test_plsql_select_into(self) -> None:
        code = "function f return char is x char; begin select a into x from meses where n = 1; return x; end;"
        self.assertEqual([span.table for span in table_spans(code)], ["MESES"])

    def test_qualified(self) -> None:
        self.assertEqual(qualified("pgpago", "bdatos", "dl_11g"), "BDATOS.PGPAGO@DL_11G")
        self.assertEqual(qualified("pgpago", "", "dl_11g"), "PGPAGO@DL_11G")
        self.assertEqual(qualified("pgpago"), "PGPAGO")


class ProjectSourcesTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        docx = build_docx(self.root / "boleta.docx", ["{{COLUMN:NRO_CUOTA}}", "{{COLUMN:F_PEN}}", "{{COLUMN:F_TOT_CUOTA}}"],
                          header="{{FIELD:F_DOBLE}}")
        result = import_docx(docx, projects_root=self.root / "proyectos", page="71", reports=FIXTURE)
        self.assertFalse(result.diagnostics.has_errors, result.diagnostics.to_dict())
        self.project = result.paths.project
        self.generated = self.project.parent

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_tables_of_queries_get_schema_and_db_link(self) -> None:
        tables = {table.table: table for table in project_tables(self.project)}
        self.assertIn("PRUEBAP_ALUMNO_CUOTA", tables)
        self.assertEqual(common_source(self.project), ("", ""))

        change = set_table_source(self.project, ["PRUEBAP_ALUMNO_CUOTA"], "bdatos", "dl_11g")
        self.assertIsNone(change.error)
        self.assertEqual(change.tables, ["PRUEBAP_ALUMNO_CUOTA"])
        query = (self.generated / "q_q_cuotas.sql").read_text(encoding="utf-8")
        self.assertIn("from BDATOS.PRUEBAP_ALUMNO_CUOTA@DL_11G c", query)
        self.assertEqual(common_source(self.project), ("BDATOS", "DL_11G"))
        self.assertTrue(validate_project(self.project).valid)

        self.assertIn("No hay cambios", set_table_source(self.project, ["PRUEBAP_ALUMNO_CUOTA"],
                                                         "bdatos", "dl_11g").error)
        removed = set_table_source(self.project, ["PRUEBAP_ALUMNO_CUOTA"])
        self.assertIsNone(removed.error)
        self.assertIn("from PRUEBAP_ALUMNO_CUOTA c", (self.generated / "q_q_cuotas.sql").read_text(encoding="utf-8"))
        self.assertIn("no es un nombre de esquema", set_table_source(self.project, None, "mal esquema", "").error)

    def test_update_docx_keeps_adjustments_and_adds_what_is_new(self) -> None:
        from docx import Document

        from corporate_report_compiler.workspace import rename_items, update_docx

        set_table_source(self.project, ["PRUEBAP_ALUMNO_CUOTA"], "BDATOS", "DL_11G")
        rename_items(self.project, {item: item.replace("P71_", "P80_") for item in
                                    [entry["item"] for entry in json.loads(self.project.read_text(encoding="utf-8"))
                                     ["parameters"]]} or {})
        before = self.project.read_text(encoding="utf-8")
        word = self.project.parent.parent / "boleta.docx"

        same = update_docx(self.project, word)                      # mismo Word: nada más cambia
        self.assertIsNotNone(same.docx)
        self.assertEqual(self.project.read_text(encoding="utf-8"), before)
        self.assertEqual(same.added_queries, [])

        document = Document(str(word))
        document.add_paragraph("Vacía: {{FIELD:NRO_CUOTA2}}")       # dato de una consulta que no estaba
        changed = self.root / "boleta_v2.docx"
        document.save(str(changed))
        update = update_docx(self.project, changed)
        self.assertEqual(update.added_queries, ["q_q_vacia.sql"])
        self.assertEqual(update.qualified_tables, ["PRUEBAP_ALUMNO_CUOTA"])
        self.assertIn("BDATOS.PRUEBAP_ALUMNO_CUOTA@DL_11G", (self.generated / "q_q_vacia.sql").read_text(encoding="utf-8"))
        raw = json.loads(self.project.read_text(encoding="utf-8"))
        self.assertIn("Q_VACIA", raw["queries"])
        self.assertTrue(all(entry["item"].startswith("P80_") for entry in raw["parameters"]))   # se conservan
        self.assertTrue((self.project.parent.parent / "boleta.docx.bak").exists() or update.backup is not None)
        self.assertTrue(validate_project(self.project).valid)

        broken = Document(str(changed))
        broken.add_paragraph("{{FIELD:mal formado!}}")
        bad = self.root / "boleta_mal.docx"
        broken.save(str(bad))
        snapshot = self.project.read_text(encoding="utf-8")
        failed = update_docx(self.project, bad)                     # Word con errores: no se toca nada
        self.assertIsNone(failed.docx)
        self.assertTrue(failed.diagnostics.has_errors)
        self.assertEqual(self.project.read_text(encoding="utf-8"), snapshot)

    def test_formula_tables_are_written_to_codigo_reports(self) -> None:
        raw = json.loads(self.project.read_text(encoding="utf-8"))
        raw["reports_code"] = "codigo_reports"
        self.project.write_text(json.dumps(raw), encoding="utf-8")
        code_dir = self.generated / "codigo_reports"
        code_dir.mkdir(exist_ok=True)
        (code_dir / "cf_dobleformula.sql").write_text(
            "function CF_DOBLEFormula return Number is x number;\n"
            "begin\n  select count(*) into x from meses;\n  return x;\nend;\n", encoding="utf-8")
        tables = {table.table: table for table in project_tables(self.project)}
        self.assertEqual(tables["MESES"].places, ["fórmula cf_dobleformula"])

        change = set_table_source(self.project, ["MESES"], "", "DL_11G")
        self.assertIsNone(change.error)
        self.assertIn("from MESES@DL_11G;", (code_dir / "cf_dobleformula.sql").read_text(encoding="utf-8"))
        result = validate_project(self.project)
        self.assertTrue(result.valid, result.diagnostics.to_dict())
        body = next(unit["body"] for unit in result.definition["reports"]["units"] if unit["name"] == "cf_dobleformula")
        self.assertIn("MESES@DL_11G", body)


if __name__ == "__main__":
    unittest.main()

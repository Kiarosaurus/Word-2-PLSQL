from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from corporate_report_compiler.diagnostics import Diagnostics
from corporate_report_compiler.docx_reader import read_template

from tests.support import DEFAULT_COLUMNS, build_valid_docx


class DocxReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def read(self, path: Path, *, strict: bool = True):
        diagnostics = Diagnostics(strict=strict)
        model = read_template(path, diagnostics)
        return diagnostics, model

    def test_reads_complete_valid_template(self) -> None:
        path = build_valid_docx(self.root / "valid.docx")
        diagnostics, model = self.read(path)

        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        assert model is not None
        self.assertEqual([column.name for column in model.columns], ["VDNI", "VNOM", "DEPARTAMENTO"])
        self.assertEqual([column.heading for column in model.columns], ["DNI", "Nombre completo", "Departamento"])
        self.assertEqual(model.fields, ("OFFICE_NAME",))
        self.assertEqual(model.template_orientation, "PORTRAIT")
        self.assertEqual(model.title_style.font_family, "HELVETICA")
        self.assertEqual(model.table_header_style.background_color, "#4A4F55")
        self.assertEqual(model.table_body_style.background_color, "#FFFFFF")
        self.assertEqual(model.footer_style.alignment, "CENTER")
        self.assertEqual(len(model.source_sha256), 64)

    def test_landscape_and_fifteen_columns_are_supported(self) -> None:
        columns = tuple((f"Columna {index}", f"COL_{index}") for index in range(1, 16))
        path = build_valid_docx(self.root / "wide.docx", columns=columns, orientation="LANDSCAPE")
        diagnostics, model = self.read(path)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        assert model is not None
        self.assertEqual(model.template_orientation, "LANDSCAPE")
        self.assertEqual(len(model.columns), 15)

    def test_more_than_fifty_columns_are_rejected(self) -> None:
        columns = tuple((f"Columna {index}", f"COL_{index}") for index in range(1, 52))
        path = build_valid_docx(self.root / "too-many-columns.docx", columns=columns)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertIn("DOCX-TABLE-010", {item.code for item in diagnostics.items})

    def test_empty_heading_is_rejected(self) -> None:
        path = build_valid_docx(self.root / "empty-heading.docx")
        document = Document(path)
        document.tables[0].cell(0, 0).text = ""
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertTrue(
            {"DOCX-TABLE-006", "DOCX-TABLE-011"} <= {item.code for item in diagnostics.items}
        )

    def test_static_header_and_footer_length_limits_are_enforced(self) -> None:
        for zone in ("header", "footer"):
            with self.subTest(zone=zone):
                path = build_valid_docx(self.root / f"long-{zone}.docx")
                document = Document(path)
                if zone == "header":
                    document.paragraphs[0].text = "{{REPORT_TITLE}}" + "X" * 4000
                else:
                    document.sections[0].footer.paragraphs[0].text = "X" * 4001
                document.save(path)
                diagnostics, model = self.read(path)
                self.assertIsNone(model)
                expected = "DOCX-STRUCT-010" if zone == "header" else "DOCX-STRUCT-011"
                self.assertIn(expected, {item.code for item in diagnostics.items})

    def test_non_a4_paper_is_rejected(self) -> None:
        path = build_valid_docx(self.root / "letter.docx")
        document = Document(path)
        section = document.sections[0]
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertIn("DOCX-STYLE-010", {item.code for item in diagnostics.items})

    def test_runtime_font_size_limits_are_enforced_by_zone(self) -> None:
        cases = (
            ("title-low", lambda doc: doc.paragraphs[0].runs, 7.5),
            ("title-high", lambda doc: doc.paragraphs[0].runs, 24.5),
            ("header-low", lambda doc: doc.tables[0].cell(0, 0).paragraphs[0].runs[-1:], 5.5),
            ("header-high", lambda doc: doc.tables[0].cell(0, 0).paragraphs[0].runs[-1:], 16.5),
            ("body-low", lambda doc: doc.tables[0].cell(1, 0).paragraphs[0].runs[-1:], 5.5),
            ("body-high", lambda doc: doc.tables[0].cell(1, 0).paragraphs[0].runs[-1:], 14.5),
            ("footer-low", lambda doc: doc.sections[0].footer.paragraphs[0].runs, 5.5),
            ("footer-high", lambda doc: doc.sections[0].footer.paragraphs[0].runs, 12.5),
        )
        for name, select_runs, size in cases:
            with self.subTest(name=name):
                path = build_valid_docx(self.root / f"{name}.docx")
                document = Document(path)
                for run in select_runs(document):
                    if run.text:
                        run.font.size = Pt(size)
                document.save(path)
                diagnostics, model = self.read(path)
                self.assertIsNone(model)
                self.assertIn("DOCX-STYLE-011", {item.code for item in diagnostics.items})

    def test_runtime_font_size_boundaries_are_accepted(self) -> None:
        path = build_valid_docx(self.root / "font-boundaries.docx")
        document = Document(path)
        for run in document.paragraphs[0].runs:
            if run.text:
                run.font.size = Pt(24)
        for cell in document.tables[0].rows[0].cells:
            cell.paragraphs[0].runs[-1].font.size = Pt(16)
        for cell in document.tables[0].rows[1].cells:
            cell.paragraphs[0].runs[-1].font.size = Pt(14)
        document.sections[0].footer.paragraphs[0].runs[0].font.size = Pt(12)
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)

    def test_body_alignment_can_vary_by_column(self) -> None:
        path = build_valid_docx(self.root / "per-column-alignment.docx")
        document = Document(path)
        document.tables[0].cell(1, 0).paragraphs[0].alignment = 1
        document.tables[0].cell(1, 2).paragraphs[0].alignment = 2
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        assert model is not None
        self.assertEqual([column.alignment for column in model.columns], ["CENTER", "START", "END"])

    def test_unspecified_body_alignment_is_preserved_for_runtime_type_inference(self) -> None:
        path = build_valid_docx(self.root / "automatic-alignment.docx")
        document = Document(path)
        document.tables[0].cell(1, 0).paragraphs[0].alignment = None
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        assert model is not None
        self.assertIsNone(model.columns[0].alignment)

    def test_word_simple_fields_and_content_controls_are_rejected(self) -> None:
        cases = (("fldSimple", "w:fldSimple"), ("sdt", "w:sdt"))
        for name, tag in cases:
            with self.subTest(name=name):
                path = build_valid_docx(self.root / f"{name}.docx")
                document = Document(path)
                element = document._element.body
                from docx.oxml import OxmlElement
                element.append(OxmlElement(tag))
                document.save(path)
                diagnostics, model = self.read(path)
                self.assertIsNone(model)
                self.assertIn("OOXML-020", {item.code for item in diagnostics.items})

    def test_heading_longer_than_runtime_limit_is_rejected(self) -> None:
        path = build_valid_docx(self.root / "long-heading.docx")
        document = Document(path)
        paragraph = document.tables[0].cell(0, 0).paragraphs[0]
        paragraph.text = "E" * 256
        paragraph.runs[0].font.name = "Arial"
        paragraph.runs[0].font.size = Pt(9)
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertIn("DOCX-TABLE-009", {item.code for item in diagnostics.items})

    def test_supported_word_fonts_map_to_native_apex_families(self) -> None:
        expected = {
            "Arial": "HELVETICA",
            "Times New Roman": "TIMES",
            "Courier New": "COURIER",
        }
        for index, (font, family) in enumerate(expected.items()):
            with self.subTest(font=font):
                path = build_valid_docx(self.root / f"font-{index}.docx", font=font)
                diagnostics, model = self.read(path)
                self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
                self.assertIsNotNone(model)
                assert model is not None
                self.assertEqual(model.title_style.font_family, family)
                self.assertEqual(model.table_header_style.font_family, family)
                self.assertEqual(model.table_body_style.font_family, family)
                self.assertEqual(model.footer_style.font_family, family)

    def test_placeholder_split_across_equally_styled_runs_is_deterministic(self) -> None:
        path = build_valid_docx(self.root / "split.docx", split_placeholder_runs=True)
        diagnostics, model = self.read(path)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        assert model is not None
        self.assertEqual(tuple(column.name for column in model.columns), tuple(alias for _, alias in DEFAULT_COLUMNS))

    def test_multiple_tables_are_rejected(self) -> None:
        path = build_valid_docx(self.root / "two-tables.docx")
        document = Document(path)
        document.add_table(rows=2, cols=1)
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertIn("DOCX-STRUCT-006", {item.code for item in diagnostics.items})

    def test_merged_cells_are_rejected(self) -> None:
        path = build_valid_docx(self.root / "merged.docx")
        document = Document(path)
        document.tables[0].cell(0, 0).merge(document.tables[0].cell(0, 1))
        document.save(path)
        diagnostics, model = self.read(path)
        self.assertIsNone(model)
        self.assertIn("DOCX-TABLE-003", {item.code for item in diagnostics.items})

    def test_unknown_font_is_an_error_in_strict_mode(self) -> None:
        path = build_valid_docx(self.root / "font.docx", font="Comic Sans MS")
        diagnostics, model = self.read(path, strict=True)
        self.assertIsNone(model)
        self.assertIn("DOCX-STYLE-002", {item.code for item in diagnostics.items})

    def test_unknown_font_is_only_a_warning_in_compatible_mode(self) -> None:
        path = build_valid_docx(self.root / "font.docx", font="Comic Sans MS")
        diagnostics, model = self.read(path, strict=False)
        self.assertFalse(diagnostics.has_errors, diagnostics.to_dict())
        self.assertIsNotNone(model)
        self.assertIn("DOCX-STYLE-002", {item.code for item in diagnostics.items})


if __name__ == "__main__":
    unittest.main()

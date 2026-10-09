"""Lectura de plantillas Word en modo layout (varias tablas, maquetación libre).

El modo layout traduce el documento completo a una descripción JSON que el
intérprete PL/SQL ``RPT_LAYOUT`` dibuja con ``RPT_PDF``:

- párrafos (bloque ``P``) con varias líneas y estilos por fragmento;
- tablas de maquetación (``G``): bordes por celda (visibles, blancos o sin
  borde), rellenos, celdas combinadas en horizontal y marcadores escalares;
- tablas de datos (``D``): la fila con ``{{COLUMN:CONSULTA.COLUMNA}}`` se repite
  por cada registro; las filas anteriores son cabecera (se repite en cada
  página) y las posteriores, totales con ``{{SUM:CONSULTA.COLUMNA}}``;
- encabezado y pie de Word, que admiten ``{{PAGE}}`` y ``{{PAGES}}``.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_UNDERLINE
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .diagnostics import Diagnostics
from .docx_reader import (
    FONT_MAP,
    MARKERS,
    _effective_bold,
    _effective_color,
    _effective_font_name,
    _effective_size,
    _scan_prohibited_elements,
    _validate_a4,
)
from .placeholders import report_empty_marker
from .security import inspect_docx_zip


TWIP_MM = 25.4 / 1440
PT_MM = 25.4 / 72
IDENT = r"[A-Z][A-Z0-9_$#]{0,29}"
TOKEN_PATTERN = re.compile(r"\{\{(.*?)\}\}", re.DOTALL)
MARKER_PATTERN = re.compile(
    rf"^(?P<kind>PAGE|PAGES|REPORT_TITLE|APP_USER|GENERATED_AT|FIELD|COLUMN|SUM)"
    rf"(?::(?P<name>{IDENT}(?:\.{IDENT})?))?(?:\|(?P<mask>[^{{}}|]{{1,60}}))?$"
)
UNQUALIFIED_PATTERN = re.compile(
    r"^\s*(?P<kind>FIELD|COLUMN|SUM)\s*:\s*(?P<name>[^|.\s][^|.]*?)\s*(?:\|(?P<mask>[^{}|]{1,60}))?$",
    re.IGNORECASE,
)
FONT_STYLE = {(False, False): 1, (True, False): 2, (False, True): 3, (True, True): 4}


@dataclass(slots=True)
class LayoutTemplate:
    source_path: Path
    source_sha256: str
    layout: dict[str, Any]
    queries: dict[str, set[str]] = field(default_factory=dict)   # consulta -> columnas usadas
    fields: set[str] = field(default_factory=set)                # {{FIELD:NOMBRE}} sin consulta
    data_tables: list[str] = field(default_factory=list)
    pending: list[str] = field(default_factory=list)             # trabajo manual detectado con Reports


def is_layout_docx(path: Path) -> bool:
    """Plantilla en modo layout: usa marcadores con CONSULTA.COLUMNA."""

    try:
        document = Document(str(path))
    except Exception:
        return False
    text = document.element.xml
    if re.search(r"\{\{\s*(COLUMN|FIELD|SUM)\s*:\s*[A-Za-z][\w$#]*\.", text) or re.search(r"\{\{\s*SUM\s*:", text):
        return True
    # El modo simple admite una sola tabla y ninguna en encabezado o pie: más que eso es
    # modo layout aunque los marcadores aún estén sin completar ({{COLUMN:}}).
    if len(document.tables) > 1:
        return True
    return any(
        bool(zone.tables)
        for section in document.sections
        for zone in (section.header, section.footer)
        if not zone.is_linked_to_previous
    )


class _Reader:
    def __init__(self, document: Any, diagnostics: Diagnostics, template: LayoutTemplate, model: Any = None) -> None:
        self.document = document
        self.diagnostics = diagnostics
        self.template = template
        self.model = model          # ReportsModel opcional: resuelve nombres de Oracle Reports
        normal = document.styles["Normal"].font
        self.default_size = float(normal.size.pt) if normal.size is not None else 11.0

    # ------------------------------------------------------------- marcadores
    def resolve_text(self, text: str, location: str) -> str:
        """Con el Modelo de Datos de Reports, completa la consulta de los marcadores.

        {{COLUMN:F_115}} o {{COLUMN:INTERES2}} -> {{COLUMN:AMORTIZACIONES.INTERES2|99,999.00}}
        """

        if self.model is None or "{{" not in text:
            return text

        def replace(match: re.Match) -> str:
            unqualified = UNQUALIFIED_PATTERN.match(match.group(1))
            if unqualified is None:
                return match.group(0)
            kind = unqualified.group("kind").upper()
            name = unqualified.group("name").strip()
            resolution = self.model.resolve(kind, name)
            if resolution is None:
                if kind != "FIELD":
                    suggestions = self.model.suggestions(name)
                    self.diagnostics.error(
                        "LAYOUT-TOKEN-009",
                        f"{match.group(0)}: {name} no existe en el Modelo de Datos de Reports.",
                        location=location,
                        suggestion=("¿Quiso decir " + ", ".join(suggestions) + "?") if suggestions else None,
                    )
                return match.group(0)
            if resolution.pending and resolution.pending not in self.template.pending:
                self.template.pending.append(resolution.pending)
                self.diagnostics.warning("LAYOUT-REPORTS-010", resolution.pending, location=location)
            mask = unqualified.group("mask") or resolution.mask
            return "{{" + f"{kind}:{resolution.marker}" + (f"|{mask}" if mask else "") + "}}"

        return TOKEN_PATTERN.sub(replace, text)

    def check_markers(self, text: str, *, zone: str, location: str, query: str | None = None) -> None:
        """zone: page (encabezado/pie), body, cell, data-header, data-body, data-footer."""

        if text.count("{{") != text.count("}}"):
            self.diagnostics.error("LAYOUT-TOKEN-001", "Hay llaves de marcador sin cerrar.", location=location)
        for match in TOKEN_PATTERN.finditer(text):
            inner = match.group(1).strip(" ").upper()
            if report_empty_marker(match.group(0), inner, self.diagnostics, location, "LAYOUT-TOKEN-008", layout=True):
                continue
            parsed = MARKER_PATTERN.match(inner)
            if parsed is None:
                self.diagnostics.error(
                    "LAYOUT-TOKEN-002", f"Marcador desconocido o mal formado: {match.group(0)}.", location=location
                )
                continue
            kind, name = parsed.group("kind"), parsed.group("name")
            needs_name = kind in {"FIELD", "COLUMN", "SUM"}
            if needs_name != (name is not None):
                self.diagnostics.error("LAYOUT-TOKEN-002", f"Marcador mal formado: {match.group(0)}.", location=location)
                continue
            if kind in {"PAGE", "PAGES"} and zone != "page":
                self.diagnostics.error(
                    "LAYOUT-TOKEN-003",
                    f"{match.group(0)} solo se admite en el encabezado o el pie de página de Word.",
                    location=location,
                )
            if kind in {"COLUMN", "SUM"}:
                if "." not in name:
                    self.diagnostics.error(
                        "LAYOUT-TOKEN-004",
                        f"{match.group(0)} debe indicar CONSULTA.COLUMNA.",
                        location=location,
                    )
                    continue
                expected = {"data-body", "data-group"} if kind == "COLUMN" else {"data-footer", "data-header", "data-group"}
                if zone not in expected:
                    where = ("la fila que se repite o en una fila de grupo de una tabla de datos" if kind == "COLUMN"
                             else "las filas que no se repiten de una tabla de datos")
                    self.diagnostics.error(
                        "LAYOUT-TOKEN-005",
                        f"{match.group(0)} solo se admite en {where}.",
                        location=location,
                    )
                    continue
                table_query, column = name.split(".", 1)
                if query is not None and table_query != query:
                    self.diagnostics.error(
                        "LAYOUT-TOKEN-006",
                        f"{match.group(0)} no pertenece a la consulta {query} de esta tabla.",
                        location=location,
                    )
                self.template.queries.setdefault(table_query, set()).add(column)
            elif kind == "FIELD":
                if "." in name:
                    source, column = name.split(".", 1)
                    self.template.queries.setdefault(source, set()).add(column)
                else:
                    self.template.fields.add(name)

    # ------------------------------------------------------------- texto
    def run_style(self, run: Any, paragraph: Paragraph, location: str) -> dict[str, Any]:
        family = _effective_font_name(run, paragraph)
        if family and FONT_MAP.get(family.casefold()) not in {None, "HELVETICA"}:
            self.diagnostics.warning(
                "LAYOUT-STYLE-001",
                f"La fuente {family} se imprime como Helvetica en el modo layout.",
                location=location,
            )
        elif family and family.casefold() not in FONT_MAP:
            self.diagnostics.warning(
                "LAYOUT-STYLE-001", f"Fuente no reconocida {family}; se imprime como Helvetica.", location=location
            )
        underline = run.underline
        if underline is None and paragraph.style is not None:
            underline = paragraph.style.font.underline
        strike = run.font.strike
        if strike is None and paragraph.style is not None:
            strike = paragraph.style.font.strike
        italic = run.italic
        if italic is None and paragraph.style is not None:
            italic = paragraph.style.font.italic
        bold = _effective_bold(run, paragraph, False)
        style = {
            "f": FONT_STYLE[(bool(bold), bool(italic))],
            "s": _effective_size(run, paragraph, self.default_size),
            "c": _effective_color(run, "#000000").lstrip("#"),
        }
        if underline not in (None, False, WD_UNDERLINE.NONE):
            style["u"] = 1                      # subrayado (cualquier tipo se imprime como línea simple)
        if strike:
            style["x"] = 1                      # tachado
        return style

    def paragraph(self, element: Any, parent: Any, *, zone: str, location: str, query: str | None = None) -> dict[str, Any]:
        paragraph = Paragraph(element, parent)
        lines: list[list[dict[str, Any]]] = [[]]
        for run in paragraph.runs:
            style = self.run_style(run, paragraph, location)
            text = run.text.replace("\t", "    ")
            for index, piece in enumerate(text.split("\n")):
                if index:
                    lines.append([])
                if not piece:
                    continue
                line = lines[-1]
                if line and all(line[-1].get(key) == style.get(key) for key in ("f", "s", "c", "u", "x")):
                    line[-1]["t"] += piece
                else:
                    line.append({"t": piece, **style})
        for line in lines:
            for run in line:
                run["t"] = self.resolve_text(run["t"], location)
                if zone == "data-group" and query:
                    # Fila de grupo: sus datos cambian con cada grupo, como una columna.
                    run["t"] = re.sub(r"\{\{\s*FIELD\s*:\s*" + re.escape(query) + r"\.", "{{COLUMN:" + query + ".",
                                      run["t"], flags=re.IGNORECASE)
                if run["t"].count("{{") != run["t"].count("}}"):
                    self.diagnostics.error(
                        "LAYOUT-TOKEN-007",
                        "Un marcador {{...}} tiene formato mixto; aplique el mismo estilo a todo el marcador.",
                        location=location,
                    )
                self.check_markers(run["t"], zone=zone, location=location, query=query)
        fmt = paragraph.paragraph_format
        style_fmt = paragraph.style.paragraph_format if paragraph.style is not None else None
        before = fmt.space_before if fmt.space_before is not None else (style_fmt.space_before if style_fmt else None)
        after = fmt.space_after if fmt.space_after is not None else (style_fmt.space_after if style_fmt else None)
        alignment = paragraph.alignment
        if alignment is None and paragraph.style is not None:
            alignment = paragraph.style.paragraph_format.alignment
        mark_size = self.default_size
        mark = element.find(qn("w:pPr"))
        if mark is not None and mark.find(qn("w:rPr")) is not None:
            size = mark.find(qn("w:rPr")).find(qn("w:sz"))
            if size is not None and size.get(qn("w:val"), "").isdigit():
                mark_size = int(size.get(qn("w:val"))) / 2
        return {
            "k": "P",
            "al": {1: "C", 2: "R", 3: "L"}.get(int(alignment) if alignment is not None else 0, "L"),
            "bf": round(before.pt * PT_MM, 2) if before is not None else 0,
            "af": round(after.pt * PT_MM, 2) if after is not None else 0,
            "es": mark_size,
            "ln": lines,
        }

    # ------------------------------------------------------------- tablas
    @staticmethod
    def _twips(element: Any, attribute: str = "w:w", default: float = 0.0) -> float:
        if element is None:
            return default
        value = element.get(qn(attribute))
        return round(int(value) * TWIP_MM, 2) if value and value.lstrip("-").isdigit() else default

    @staticmethod
    def _border(element: Any) -> list[Any] | None:
        if element is None or element.get(qn("w:val"), "nil") in {"nil", "none"}:
            return None
        size = element.get(qn("w:sz"))
        color = element.get(qn("w:color"))
        width = round(int(size) / 8, 2) if size and size.isdigit() else 0.5
        color = color.upper() if color and color.casefold() != "auto" else "000000"
        return [max(width, 0.25), color]

    def _table_borders(self, tbl: Any) -> Any:
        props = tbl.find(qn("w:tblPr"))
        borders = props.find(qn("w:tblBorders")) if props is not None else None
        if borders is not None:
            return borders
        style = props.find(qn("w:tblStyle")) if props is not None else None
        if style is not None:
            for candidate in self.document.styles.element.iter(qn("w:style")):
                if candidate.get(qn("w:styleId")) == style.get(qn("w:val")):
                    style_props = candidate.find(qn("w:tblPr"))
                    if style_props is not None:
                        return style_props.find(qn("w:tblBorders"))
        return None

    def table(self, tbl: Any, parent: Any, *, zone: str, content_width: float, location: str) -> dict[str, Any] | None:
        props = tbl.find(qn("w:tblPr"))
        widths = [self._twips(col) for col in tbl.find(qn("w:tblGrid")).iter(qn("w:gridCol"))]
        total = sum(widths)
        x = 0.0
        if props is not None:
            jc = props.find(qn("w:jc"))
            align = jc.get(qn("w:val")) if jc is not None else None
            if align == "center":
                x = round((content_width - total) / 2, 2)
            elif align in {"right", "end"}:
                x = round(content_width - total, 2)
            else:
                x = self._twips(props.find(qn("w:tblInd")))
        margins = props.find(qn("w:tblCellMar")) if props is not None else None

        def margin(side: str, alt: str, default: float) -> float:
            if margins is None:
                return default
            node = margins.find(qn(f"w:{side}"))
            if node is None:
                node = margins.find(qn(f"w:{alt}"))
            return self._twips(node, default=default)

        pad = [margin("left", "start", 1.9), margin("top", "top", 0.0), margin("right", "end", 1.9), margin("bottom", "bottom", 0.0)]
        table_borders = self._table_borders(tbl)

        def table_side(name: str) -> list[Any] | None:
            return self._border(table_borders.find(qn(f"w:{name}"))) if table_borders is not None else None

        rows_xml = list(tbl.iter(qn("w:tr")))
        if any(nested is not tbl for nested in tbl.iter(qn("w:tbl"))):
            self.diagnostics.error("LAYOUT-TABLE-001", "No se admiten tablas anidadas.", location=location)
            return None
        text = "".join(t.text or "" for t in tbl.iter(qn("w:t")))
        is_data = "{{COLUMN:" in re.sub(r"\s", "", text).upper()
        plan = self._data_sections(rows_xml, location) if is_data else None
        if is_data and plan is None:
            return None

        rows: list[dict[str, Any]] = []
        last_row = len(rows_xml) - 1
        for row_index, row in enumerate(rows_xml):
            row_location = f"{location}, fila {row_index + 1}"
            query = None
            if is_data:
                section, row_zone = plan["rows"][row_index]
                query = plan["queries"][section]
            else:
                row_zone = zone
            tr_props = row.find(qn("w:trPr"))
            height = tr_props.find(qn("w:trHeight")) if tr_props is not None else None
            if tr_props is not None and tr_props.find(qn("w:gridBefore")) is not None:
                self.diagnostics.error("LAYOUT-TABLE-003", "No se admiten filas desplazadas (gridBefore).", location=row_location)
            cells: list[dict[str, Any]] = []
            column = 0
            cells_xml = [child for child in row if child.tag == qn("w:tc")]
            for cell in cells_xml:
                cell_location = f"{row_location}, columna {column + 1}"
                tc_props = cell.find(qn("w:tcPr"))
                span_node = tc_props.find(qn("w:gridSpan")) if tc_props is not None else None
                span = int(span_node.get(qn("w:val"), "1")) if span_node is not None else 1
                if tc_props is not None and tc_props.find(qn("w:vMerge")) is not None:
                    self.diagnostics.error(
                        "LAYOUT-TABLE-004",
                        "No se admiten celdas combinadas en vertical; combine solo en horizontal.",
                        location=cell_location,
                    )
                cell_borders = tc_props.find(qn("w:tcBorders")) if tc_props is not None else None

                def cell_side(names: tuple[str, ...], fallback: str) -> list[Any] | None:
                    if cell_borders is not None:
                        for name in names:
                            node = cell_borders.find(qn(f"w:{name}"))
                            if node is not None:
                                return self._border(node)
                    return table_side(fallback)

                borders = {
                    "t": cell_side(("top",), "top" if row_index == 0 else "insideH"),
                    "b": cell_side(("bottom",), "bottom" if row_index == last_row else "insideH"),
                    "l": cell_side(("left", "start"), "left" if column == 0 else "insideV"),
                    "r": cell_side(("right", "end"), "right" if column + span >= len(widths) else "insideV"),
                }
                shading = tc_props.find(qn("w:shd")) if tc_props is not None else None
                fill = shading.get(qn("w:fill")) if shading is not None else None
                fill = fill.upper() if fill and fill.casefold() != "auto" else None
                valign = tc_props.find(qn("w:vAlign")) if tc_props is not None else None
                valign = {"center": "C", "bottom": "B"}.get(valign.get(qn("w:val")) if valign is not None else "", "T")
                paragraphs = [
                    self.paragraph(p, parent, zone=row_zone if is_data else ("page" if zone == "page" else "cell"),
                                   location=cell_location, query=query)
                    for p in cell if p.tag == qn("w:p")
                ]
                cells.append({"c": column, "sp": span, "fl": fill, "va": valign,
                              "b": {key: value for key, value in borders.items() if value}, "ps": paragraphs})
                column += span
            if column != len(widths):
                self.diagnostics.error(
                    "LAYOUT-TABLE-005",
                    "La fila no ocupa todas las columnas de la tabla.",
                    location=row_location,
                )
            rule = height.get(qn("w:hRule"), "atLeast") if height is not None else "auto"
            rows.append({"h": self._twips(height, "w:val"), "hr": rule if rule in {"exact", "atLeast"} else "auto", "cs": cells})

        block = {"x": x, "w": widths, "pd": pad}
        if not is_data:
            return {"k": "G", **block, "rs": rows}
        if zone == "page":
            self.diagnostics.error(
                "LAYOUT-TABLE-006", "Las tablas de datos no se admiten en el encabezado o el pie.", location=location
            )
            return None

        def sums_in(selected: list[dict[str, Any]]) -> set[str]:
            return {
                match.group(2)
                for row in selected
                for cell in row["cs"]
                for para in cell["ps"]
                for line in para["ln"]
                for run in line
                for match in re.finditer(r"\{\{\s*SUM\s*:\s*(" + IDENT + r")\.(" + IDENT + r")", run["t"].upper())
            }

        # Una tabla de Word puede tener varias secciones de datos (una por consulta), una
        # debajo de otra, como los marcos de Oracle Reports: cada una se dibuja como tabla propia.
        blocks = []
        for section, query in enumerate(plan["queries"]):
            indexes = [i for i, (owner, _zone) in enumerate(plan["rows"]) if owner == section]
            header = [rows[i] for i in indexes if plan["rows"][i][1] in {"data-header", "data-group"}]
            for i in indexes:
                if plan["rows"][i][1] == "data-group":
                    rows[i]["g"] = plan["groups"][i]
            body = next(rows[i] for i in indexes if plan["rows"][i][1] == "data-body")
            footer = [rows[i] for i in indexes if plan["rows"][i][1] == "data-footer"]
            data = {"k": "D", "q": query, **block, "hd": header, "bd": body, "ft": footer,
                    "sm": sorted(sums_in(header + footer))}
            if sums_in([rows[i] for i in indexes if plan["rows"][i][1] == "data-header"]):
                data["hs"] = True              # totales arriba de las filas: se calculan antes de dibujarlas
            if self.model is not None:
                data["he"] = True              # como un marco de Reports: sin filas, la sección no se imprime
            self.template.data_tables.append(query)
            blocks.append(data)
        return blocks

    def _marker_refs(self, text: str) -> list[tuple[str, str | None, str | None]]:
        """(tipo, consulta, columna) de cada marcador FIELD/COLUMN/SUM del texto, ya resuelto con Reports."""

        refs = []
        for match in TOKEN_PATTERN.finditer(text):
            parsed = re.match(r"^\s*(FIELD|COLUMN|SUM)\s*:\s*([^|]+?)\s*(?:\|.*)?$", match.group(1), re.IGNORECASE | re.DOTALL)
            if parsed is None:
                continue
            kind, name = parsed.group(1).upper(), parsed.group(2).strip()
            if "." in name:
                query, column = name.upper().split(".", 1)
            elif self.model is not None and (resolution := self.model.resolve(kind, name)) and "." in resolution.marker:
                query, column = resolution.marker.split(".", 1)
            else:
                query = column = None
            refs.append((kind, query, column))
        return refs

    def _data_sections(self, rows_xml: list[Any], location: str) -> dict[str, Any] | None:
        """Reparte las filas de una tabla de datos en secciones (una por fila {{COLUMN:...}}).

        Cada fila que no se repite va con la sección cuya consulta usa; si no usa ninguna,
        con la sección siguiente (título o encabezados) o, al final, con la última (totales).
        Dentro de una sección, las filas de encabezado que muestran datos de su consulta son
        filas de grupo: se repiten cada vez que cambian esos datos (grupos de Reports).
        """

        refs = [self._marker_refs("".join(t.text or "" for t in row.iter(qn("w:t")))) for row in rows_xml]
        bodies: list[int] = []
        queries: list[str | None] = []
        for index, row in enumerate(rows_xml):
            row_text = re.sub(r"\s", "", "".join(t.text or "" for t in row.iter(qn("w:t")))).upper()
            if "{{COLUMN:" not in row_text:
                continue
            # Consulta de la sección: la de su primer {{COLUMN:}} conocido. Si la fila mezcla
            # consultas, la revisión de marcadores lo informa (LAYOUT-TOKEN-006).
            bodies.append(index)
            queries.append(next((query for kind, query, _c in refs[index] if kind == "COLUMN" and query), None))
        assignment: list[tuple[int, str]] = []
        for index, row_refs in enumerate(refs):
            if index in bodies:
                assignment.append((bodies.index(index), "data-body"))
                continue
            following = next((k for k, body in enumerate(bodies) if body > index), None)
            previous = following - 1 if following is not None else len(bodies) - 1
            used = {query for _kind, query, _column in row_refs}
            if following is None:
                section = previous
            elif previous >= 0 and queries[previous] in used and queries[following] not in used:
                section = previous
            else:
                section = following
            assignment.append((section, "data-header" if index < bodies[section] else "data-footer"))
        for index in range(1, len(assignment)):
            if assignment[index][0] < assignment[index - 1][0]:
                self.diagnostics.error(
                    "LAYOUT-TABLE-002",
                    "Las filas de cada consulta deben ir juntas: totales debajo de su fila {{COLUMN:...}} "
                    "y títulos o encabezados de la siguiente después.",
                    location=f"{location}, fila {index + 1}",
                )
                return None
        groups: dict[int, list[str]] = {}
        for section, query in enumerate(queries):
            keys: list[str] = []
            header = [i for i, (owner, zone) in enumerate(assignment) if owner == section and zone == "data-header"]
            seen_group = False
            for index in header:
                own = [column for kind, q, column in refs[index] if kind in {"FIELD", "COLUMN"} and q == query and column]
                if own:
                    seen_group = True
                    keys += [column for column in own if column not in keys]
                    assignment[index] = (section, "data-group")
                    groups[index] = list(keys)
                elif seen_group:
                    self.diagnostics.error(
                        "LAYOUT-TABLE-007",
                        "Las filas de grupo (con datos de la consulta de la tabla) deben ir justo antes de la fila "
                        "que se repite.",
                        location=f"{location}, fila {index + 1}",
                    )
                    return None
        return {"queries": queries, "rows": assignment, "groups": groups}

    # ------------------------------------------------------------- zonas
    def blocks(self, container: Any, parent: Any, *, zone: str, content_width: float, label: str) -> list[dict[str, Any]]:
        result: list[dict[str, Any]] = []
        tables = 0
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                result.append(self.paragraph(child, parent, zone=zone, location=f"{label}, párrafo {len(result) + 1}"))
            elif child.tag == qn("w:tbl"):
                tables += 1
                block = self.table(child, parent, zone=zone, content_width=content_width, location=f"{label}, tabla {tables}")
                if isinstance(block, list):
                    result.extend(block)
                elif block is not None:
                    result.append(block)
            elif child.tag == qn("w:sectPr") or child.tag in MARKERS:
                continue
            else:
                self.diagnostics.error(
                    "LAYOUT-STRUCT-001",
                    "Estructura de Word no admitida.",
                    location=f"{label}: {str(child.tag).rsplit('}', 1)[-1]}",
                )
        # Párrafos vacíos al final del encabezado/pie no ocupan espacio útil.
        while zone == "page" and result and result[-1]["k"] == "P" and not any(result[-1]["ln"]):
            result.pop()
        return result


def read_layout_template(path: Path, diagnostics: Diagnostics, model: Any = None) -> LayoutTemplate | None:
    if inspect_docx_zip(path, diagnostics) is None:
        return None
    try:
        document = Document(str(path))
    except Exception as exc:
        diagnostics.error("DOCX-STRUCT-001", f"Word no pudo interpretar la plantilla: {exc}", location=str(path))
        return None
    _scan_prohibited_elements(document, diagnostics)
    if len(document.sections) != 1:
        diagnostics.error("DOCX-STRUCT-002", "La plantilla debe contener una sola sección de Word.")
        return None
    section = document.sections[0]
    if section.different_first_page_header_footer:
        diagnostics.error("DOCX-STRUCT-003", "No se admite un encabezado o pie diferente en la primera página.")
    _validate_a4(section, diagnostics)
    if diagnostics.has_errors:
        return None

    width = round(section.page_width.mm, 2)
    height = round(section.page_height.mm, 2)
    page = {
        "w": width,
        "h": height,
        "mt": round(section.top_margin.mm, 2),
        "mb": round(section.bottom_margin.mm, 2),
        "ml": round(section.left_margin.mm, 2),
        "mr": round(section.right_margin.mm, 2),
        "hd": round(section.header_distance.mm, 2),
        "fd": round(section.footer_distance.mm, 2),
    }
    content_width = width - page["ml"] - page["mr"]
    template = LayoutTemplate(path, sha256(path.read_bytes()).hexdigest(), {})
    reader = _Reader(document, diagnostics, template, model)
    header = [] if section.header.is_linked_to_previous else reader.blocks(
        section.header._element, section.header, zone="page", content_width=content_width, label="Encabezado")
    footer = [] if section.footer.is_linked_to_previous else reader.blocks(
        section.footer._element, section.footer, zone="page", content_width=content_width, label="Pie")
    body = reader.blocks(document.element.body, document._body, zone="body", content_width=content_width, label="Cuerpo")
    if not template.data_tables:
        diagnostics.warning("LAYOUT-STRUCT-002", "La plantilla no tiene ninguna tabla de datos {{COLUMN:...}}.")
    if diagnostics.has_errors:
        return None
    template.layout = {"page": page, "header": header, "body": body, "footer": footer}
    return template

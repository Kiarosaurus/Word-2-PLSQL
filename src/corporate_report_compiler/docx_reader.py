from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .diagnostics import Diagnostics
from .model import ColumnModel, TemplateModel, TextStyle
from .placeholders import (
    parse_column_placeholder,
    validate_text_placeholders,
)
from .security import inspect_docx_zip


FONT_MAP = {
    "arial": "HELVETICA",
    "aptos": "HELVETICA",
    "aptos display": "HELVETICA",
    "aptos narrow": "HELVETICA",
    "calibri": "HELVETICA",
    "helvetica": "HELVETICA",
    "liberation sans": "HELVETICA",
    "times new roman": "TIMES",
    "cambria": "TIMES",
    "georgia": "TIMES",
    "liberation serif": "TIMES",
    "courier new": "COURIER",
    "consolas": "COURIER",
    "liberation mono": "COURIER",
}

PROHIBITED_TAGS = {
    "drawing": "imagen, forma o gráfico",
    "pict": "imagen o forma heredada",
    "object": "objeto OLE",
    "altChunk": "contenido externo insertado",
    "fldChar": "campo de Word",
    "instrText": "instrucción de campo de Word",
    "fldSimple": "campo simple de Word",
    "sdt": "control de contenido de Word",
    "ins": "cambio controlado pendiente",
    "del": "cambio controlado pendiente",
    "moveFrom": "cambio controlado pendiente",
    "moveTo": "cambio controlado pendiente",
    "txbxContent": "cuadro de texto",
    # Envoltorios cuyo texto Word muestra pero python-docx no expone en
    # ``Paragraph.text``: aceptarlos haría que la salida difiera del diseño.
    "smartTag": "etiqueta inteligente de Word",
    "customXml": "XML personalizado en línea",
    "dir": "bloque de dirección de texto",
    "bdo": "bloque de dirección de texto",
    "ruby": "texto ruby",
    # Revisiones de formato pendientes.
    "rPrChange": "cambio controlado pendiente",
    "pPrChange": "cambio controlado pendiente",
    "sectPrChange": "cambio controlado pendiente",
    "tblPrChange": "cambio controlado pendiente",
    "tblPrExChange": "cambio controlado pendiente",
    "trPrChange": "cambio controlado pendiente",
    "tcPrChange": "cambio controlado pendiente",
    "tblGridChange": "cambio controlado pendiente",
    "numberingChange": "cambio controlado pendiente",
    "cellIns": "cambio controlado pendiente",
    "cellDel": "cambio controlado pendiente",
    "cellMerge": "cambio controlado pendiente",
    "moveFromRangeStart": "cambio controlado pendiente",
    "moveToRangeStart": "cambio controlado pendiente",
    # Notas y comentarios.
    "footnoteReference": "nota al pie",
    "endnoteReference": "nota al final",
    "commentReference": "comentario",
    "commentRangeStart": "comentario",
    "commentRangeEnd": "comentario",
    # Texto oculto: Word no lo muestra, pero llegaría al reporte APEX.
    "vanish": "texto oculto",
    "specVanish": "texto oculto",
    "webHidden": "texto oculto",
    # Contenido visible en Word que python-docx no incluye en el texto.
    "sym": "símbolo insertado",
    "pgNum": "número de página de Word",
    "subDoc": "subdocumento",
}
HIDDEN_TEXT_TAGS = {"vanish", "specVanish", "webHidden"}
FALSE_VALUES = {"0", "false", "off"}
MC_ALTERNATE_CONTENT = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
MATH_NAMESPACE = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MARKERS = {
    qn("w:bookmarkStart"),
    qn("w:bookmarkEnd"),
    qn("w:proofErr"),
    qn("w:permStart"),
    qn("w:permEnd"),
}
ALLOWED_BODY_CHILDREN = {qn("w:p"), qn("w:tbl"), qn("w:sectPr"), *MARKERS}
ALLOWED_HEADER_FOOTER_CHILDREN = {qn("w:p"), *MARKERS}


def _utf8_length(value: str) -> int:
    return len(value.encode("utf-8"))


def _alignment(value: object, default: str | None = "START") -> str | None:
    if value == WD_ALIGN_PARAGRAPH.CENTER:
        return "CENTER"
    if value == WD_ALIGN_PARAGRAPH.RIGHT:
        return "END"
    if value == WD_ALIGN_PARAGRAPH.LEFT:
        return "START"
    return default


def _effective_font_name(run: Run, paragraph: Paragraph) -> str | None:
    candidates = [
        run.font.name,
        getattr(getattr(run, "style", None), "font", None).name if getattr(run, "style", None) else None,
        paragraph.style.font.name if paragraph.style is not None else None,
    ]
    for candidate in candidates:
        if candidate:
            return str(candidate)
    return None


def _effective_size(run: Run, paragraph: Paragraph, default: float) -> float:
    value = run.font.size
    if value is None and getattr(run, "style", None) is not None:
        value = run.style.font.size
    if value is None and paragraph.style is not None:
        value = paragraph.style.font.size
    return round(float(value.pt), 2) if value is not None else default


def _effective_bold(run: Run, paragraph: Paragraph, default: bool) -> bool:
    value = run.bold
    if value is None and getattr(run, "style", None) is not None:
        value = run.style.font.bold
    if value is None and paragraph.style is not None:
        value = paragraph.style.font.bold
    return default if value is None else bool(value)


def _effective_color(run: Run, default: str) -> str:
    rgb = run.font.color.rgb
    return f"#{rgb}".upper() if rgb is not None else default


def _run_style(
    run: Run,
    paragraph: Paragraph,
    diagnostics: Diagnostics,
    *,
    location: str,
    default_size: float,
    default_color: str,
    default_bold: bool,
) -> tuple[str, float, str, str]:
    word_font = _effective_font_name(run, paragraph)
    if word_font is None:
        diagnostics.warning(
            "DOCX-STYLE-001",
            "No se pudo determinar la fuente efectiva; se usará Helvetica.",
            location=location,
            strict_error=False,
        )
        family = "HELVETICA"
    else:
        family = FONT_MAP.get(word_font.casefold(), "")
        if not family:
            diagnostics.warning(
                "DOCX-STYLE-002",
                f"La fuente {word_font!r} no es compatible; se usará Helvetica.",
                location=location,
                suggestion="Use Arial, Aptos, Calibri, Times New Roman, Cambria, Georgia, Courier New o Consolas.",
                strict_error=True,
            )
            family = "HELVETICA"
    for attr, label in ((run.italic, "cursiva"), (run.underline, "subrayado"), (run.font.strike, "tachado")):
        if attr:
            diagnostics.warning(
                "DOCX-STYLE-003",
                f"El formato {label} no puede reproducirse y se ignorará.",
                location=location,
                strict_error=True,
            )
    return (
        family,
        _effective_size(run, paragraph, default_size),
        "BOLD" if _effective_bold(run, paragraph, default_bold) else "NORMAL",
        _effective_color(run, default_color),
    )


def _paragraph_style(
    paragraph: Paragraph,
    diagnostics: Diagnostics,
    *,
    location: str,
    default_size: float,
    default_color: str,
    default_bold: bool,
    minimum_size: float,
    maximum_size: float,
    forced_alignment: str | None = None,
    default_alignment: str | None = "START",
) -> TextStyle:
    runs = [run for run in paragraph.runs if run.text and not run.text.isspace()]
    if not runs:
        values = ("HELVETICA", default_size, "BOLD" if default_bold else "NORMAL", default_color)
    else:
        styles = {
            _run_style(
                run,
                paragraph,
                diagnostics,
                location=location,
                default_size=default_size,
                default_color=default_color,
                default_bold=default_bold,
            )
            for run in runs
        }
        values = next(iter(styles))
        if len(styles) > 1:
            diagnostics.warning(
                "DOCX-STYLE-004",
                "La zona contiene formato de texto mixto; APEX requiere un estilo uniforme.",
                location=location,
                suggestion="Seleccione todo el texto de la zona y aplique una sola fuente, tamaño, peso y color.",
                strict_error=True,
            )
    if not minimum_size <= values[1] <= maximum_size:
        diagnostics.error(
            "DOCX-STYLE-011",
            f"El tamaño de fuente debe estar entre {minimum_size:g} y {maximum_size:g} puntos.",
            location=location,
        )
    return TextStyle(
        font_family=values[0],
        font_size=values[1],
        font_weight=values[2],
        font_color=values[3],
        alignment=forced_alignment or _alignment(paragraph.alignment, default_alignment),
    )


def _cell_paragraph(cell: _Cell, diagnostics: Diagnostics, location: str) -> Paragraph:
    nonempty = [paragraph for paragraph in cell.paragraphs if paragraph.text.strip()]
    if len(nonempty) != 1:
        diagnostics.error(
            "DOCX-TABLE-006",
            "Cada celda debe contener exactamente un párrafo con contenido.",
            location=location,
        )
    return nonempty[0] if nonempty else cell.paragraphs[0]


def _cell_fill(cell: _Cell, default: str, diagnostics: Diagnostics, location: str) -> str:
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return default
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        return default
    if any(shading.get(qn(attribute)) for attribute in ("w:themeFill", "w:themeFillTint", "w:themeFillShade")):
        diagnostics.error(
            "DOCX-STYLE-012",
            "Los colores de tema no son compatibles; aplique un color RGB explícito al relleno.",
            location=location,
        )
    value = shading.get(qn("w:fill"))
    if not value or value.casefold() in {"auto", "none"}:
        return default
    return f"#{value[-6:]}".upper()


def _border_values(borders: object, diagnostics: Diagnostics, location: str) -> list[tuple[float, str]]:
    values: list[tuple[float, str]] = []
    if borders is None:
        return values
    for side in ("insideH", "insideV", "top", "bottom", "left", "right"):
        element = borders.find(qn(f"w:{side}"))
        if element is None or element.get(qn("w:val"), "nil") in {"nil", "none"}:
            continue
        if element.get(qn("w:themeColor")):
            diagnostics.error(
                "DOCX-STYLE-013",
                "Los colores de tema no son compatibles en bordes; aplique un color RGB explícito.",
                location=location,
            )
        size = element.get(qn("w:sz"))
        color = element.get(qn("w:color"))
        width = round(float(size) / 8, 2) if size and size.isdigit() else 0.5
        normalized = f"#{color[-6:]}".upper() if color and color.casefold() != "auto" else "#BFC3C7"
        values.append((width, normalized))
    return values


def _table_border(table: object, diagnostics: Diagnostics) -> tuple[float, str]:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
    values = _border_values(borders, diagnostics, "Tabla principal")
    for row_index, row in enumerate(table.rows, start=1):
        for cell_index, cell in enumerate(row.cells, start=1):
            tc_pr = cell._tc.tcPr
            cell_borders = tc_pr.find(qn("w:tcBorders")) if tc_pr is not None else None
            values.extend(
                _border_values(
                    cell_borders,
                    diagnostics,
                    f"Tabla 1, fila {row_index}, celda {cell_index}",
                )
            )
    distinct = set(values)
    if len(distinct) > 1:
        diagnostics.error(
            "DOCX-STYLE-014",
            "Todos los bordes de la tabla deben usar el mismo grosor y color RGB.",
            location="Tabla principal",
        )
    return values[0] if values else (0.5, "#BFC3C7")


def _column_weights(table: object, count: int) -> list[float]:
    grid = table._tbl.tblGrid
    widths: list[float] = []
    if grid is not None:
        for child in grid:
            value = child.get(qn("w:w"))
            if value:
                try:
                    widths.append(float(value))
                except ValueError:
                    pass
    if len(widths) != count:
        widths = []
        for cell in table.rows[0].cells:
            value = getattr(cell, "width", None)
            widths.append(float(value or 1))
    minimum = min((value for value in widths if value > 0), default=1.0)
    return [round(max(value, minimum) / minimum, 4) for value in widths]


def _validate_a4(section: object, diagnostics: Diagnostics) -> None:
    if section.page_width is None or section.page_height is None:
        diagnostics.error(
            "DOCX-STYLE-010",
            "La plantilla debe declarar un tamaño de página A4.",
            location="Configuración de página",
        )
        return
    width = float(section.page_width.inches)
    height = float(section.page_height.inches)
    short, long = sorted((width, height))
    distance = abs(short - 8.27) + abs(long - 11.69)
    if distance > 0.3:
        diagnostics.error(
            "DOCX-STYLE-010",
            "La plantilla debe usar papel A4; el tamaño no se configura en el proyecto.",
            location="Configuración de página",
        )


def _scan_prohibited_elements(document: object, diagnostics: Diagnostics) -> None:
    roots: list[object] = [document.element]
    for section in document.sections:
        roots.extend([section.header._element, section.footer._element])
    # Un estilo con texto oculto lo aplica a todo párrafo que lo use.
    roots.append(document.styles.element)
    reported: set[str] = set()
    for root in roots:
        for local_name, label in PROHIBITED_TAGS.items():
            if label in reported:
                continue
            elements = root.iter(qn(f"w:{local_name}"))
            if local_name in HIDDEN_TEXT_TAGS:
                # <w:vanish w:val="0"/> desactiva el texto oculto heredado.
                elements = (
                    element for element in elements
                    if (element.get(qn("w:val")) or "true").casefold() not in FALSE_VALUES
                )
            if any(True for _ in elements):
                reported.add(label)
                diagnostics.error("OOXML-020", f"La plantilla contiene {label}, que no está permitido.")
        if "ecuación" not in reported and any(
            str(element.tag).startswith(MATH_NAMESPACE) for element in root.iter()
        ):
            reported.add("ecuación")
            diagnostics.error("OOXML-020", "La plantilla contiene una ecuación, que no está permitida.")
        if "contenido alternativo" not in reported and any(True for _ in root.iter(MC_ALTERNATE_CONTENT)):
            reported.add("contenido alternativo")
            diagnostics.error(
                "OOXML-020",
                "La plantilla contiene contenido alternativo de compatibilidad, que no está permitido.",
            )


def read_template(path: Path, diagnostics: Diagnostics) -> TemplateModel | None:
    if inspect_docx_zip(path, diagnostics) is None:
        return None
    digest = sha256(path.read_bytes()).hexdigest()
    try:
        document = Document(str(path))
    except Exception as exc:
        diagnostics.error("DOCX-STRUCT-001", f"Word no pudo interpretar la plantilla: {exc}", location=str(path))
        return None

    try:
        _scan_prohibited_elements(document, diagnostics)
        settings = document.settings.element
    except AttributeError:
        diagnostics.error(
            "OOXML-003",
            "Las partes de estilos o configuración del documento no tienen el tipo esperado.",
            location=str(path),
        )
        return None
    if len(document.sections) != 1:
        diagnostics.error("DOCX-STRUCT-002", "La plantilla debe contener una sola sección de Word.")
        return None
    section = document.sections[0]
    if section.different_first_page_header_footer:
        diagnostics.error("DOCX-STRUCT-003", "No se admite un encabezado o footer diferente en la primera página.")
    if settings.find(qn("w:evenAndOddHeaders")) is not None:
        diagnostics.error("DOCX-STRUCT-004", "No se admiten encabezados o footers diferentes para páginas pares.")
    for label, part in (("encabezado", section.header), ("pie", section.footer)):
        for child in part._element.iterchildren():
            if child.tag not in ALLOWED_HEADER_FOOTER_CHILDREN:
                diagnostics.error(
                    "DOCX-STRUCT-012",
                    f"El {label} de Word contiene una estructura no admitida (por ejemplo, una tabla).",
                    location=str(child.tag).rsplit("}", 1)[-1],
                )
                return None
    if any(paragraph.text.strip() for paragraph in section.header.paragraphs):
        diagnostics.error(
            "DOCX-STRUCT-005",
            "Use el párrafo anterior a la tabla para el encabezado; el encabezado real de Word debe estar vacío.",
        )

    for child in document.element.body.iterchildren():
        if child.tag not in ALLOWED_BODY_CHILDREN:
            diagnostics.error(
                "DOCX-STRUCT-012",
                "El cuerpo del documento contiene una estructura no admitida.",
                location=str(child.tag).rsplit("}", 1)[-1],
            )
            return None

    body_items: list[tuple[str, object]] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            body_items.append(("P", Paragraph(child, document._body)))
        elif child.tag == qn("w:tbl"):
            matching = next((table for table in document.tables if table._tbl is child), None)
            body_items.append(("T", matching))
    tables = [item for kind, item in body_items if kind == "T" and item is not None]
    if len(tables) != 1:
        diagnostics.error("DOCX-STRUCT-006", "La plantilla debe contener exactamente una tabla principal.")
        return None
    table = tables[0]
    table_index = next(index for index, item in enumerate(body_items) if item[1] is table)
    before = [item for kind, item in body_items[:table_index] if kind == "P" and item.text.strip()]
    after = [item for kind, item in body_items[table_index + 1 :] if kind == "P" and item.text.strip()]
    if len(before) != 1:
        diagnostics.error(
            "DOCX-STRUCT-007",
            "Debe existir un único párrafo de encabezado con contenido antes de la tabla.",
        )
        return None
    if after:
        diagnostics.error("DOCX-STRUCT-008", "No se permite contenido libre después de la tabla.")
    header = before[0]
    if _utf8_length(header.text) > 4000:
        diagnostics.error(
            "DOCX-STRUCT-010",
            "El encabezado no puede superar 4000 bytes UTF-8 antes de sustituir campos.",
        )
    if header.text.upper().count("{{REPORT_TITLE}}") != 1:
        diagnostics.error(
            "TOKEN-010",
            "El encabezado debe contener {{REPORT_TITLE}} exactamente una vez.",
            location="Párrafo de encabezado",
        )
    header_fields = validate_text_placeholders(
        header.text,
        allowed_kinds={"REPORT_TITLE", "APP_USER", "GENERATED_AT", "FIELD"},
        diagnostics=diagnostics,
        location="Párrafo de encabezado",
    )

    if len(table.rows) != 2:
        diagnostics.error("DOCX-TABLE-001", "La tabla principal debe tener exactamente dos filas.")
        return None
    header_cells = table.rows[0].cells
    body_cells = table.rows[1].cells
    if len(header_cells) != len(body_cells) or not header_cells:
        diagnostics.error("DOCX-TABLE-002", "Las dos filas deben contener la misma cantidad de celdas.")
        return None
    if len(header_cells) > 50:
        diagnostics.error(
            "DOCX-TABLE-010",
            "La tabla no puede contener más de 50 columnas.",
            location="Tabla principal",
        )
    for element in table._tbl.iter():
        if element.tag in {qn("w:gridSpan"), qn("w:vMerge")}:
            diagnostics.error("DOCX-TABLE-003", "No se permiten celdas combinadas.", location="Tabla principal")
    for row_index, row in enumerate(table.rows, start=1):
        for cell_index, cell in enumerate(row.cells, start=1):
            if cell.tables:
                diagnostics.error(
                    "DOCX-TABLE-005",
                    "No se permiten tablas anidadas.",
                    location=f"Tabla 1, fila {row_index}, celda {cell_index}",
                )

    columns: list[ColumnModel] = []
    names: set[str] = set()
    weights = _column_weights(table, len(header_cells))
    header_styles: list[TextStyle] = []
    body_styles: list[TextStyle] = []
    header_fills: list[str] = []
    body_fills: list[str] = []
    for index, (header_cell, body_cell) in enumerate(zip(header_cells, body_cells), start=1):
        header_location = f"Tabla 1, fila 1, celda {index}"
        body_location = f"Tabla 1, fila 2, celda {index}"
        header_paragraph = _cell_paragraph(header_cell, diagnostics, header_location)
        body_paragraph = _cell_paragraph(body_cell, diagnostics, body_location)
        if not header_paragraph.text.strip():
            diagnostics.error(
                "DOCX-TABLE-011",
                "La etiqueta visible del encabezado no puede estar vacía.",
                location=header_location,
            )
        if "{{" in header_paragraph.text or "}}" in header_paragraph.text:
            diagnostics.error("DOCX-TABLE-007", "La etiqueta del encabezado debe ser texto literal.", location=header_location)
        name = parse_column_placeholder(body_paragraph.text, diagnostics, body_location)
        if name is None:
            continue
        if name in names:
            diagnostics.error("DOCX-TABLE-008", f"La columna {name} está duplicada.", location=body_location)
            continue
        names.add(name)
        heading = header_paragraph.text.strip()
        if _utf8_length(heading) > 255:
            diagnostics.error(
                "DOCX-TABLE-009",
                "La etiqueta del encabezado no puede superar 255 bytes UTF-8.",
                location=header_location,
            )
        header_style = _paragraph_style(
            header_paragraph,
            diagnostics,
            location=header_location,
            default_size=9,
            default_color="#FFFFFF",
            default_bold=True,
            minimum_size=6,
            maximum_size=16,
        )
        body_style = _paragraph_style(
            body_paragraph,
            diagnostics,
            location=body_location,
            default_size=8.5,
            default_color="#25282B",
            default_bold=False,
            minimum_size=6,
            maximum_size=14,
            default_alignment=None,
        )
        header_styles.append(header_style)
        body_styles.append(body_style)
        header_fills.append(_cell_fill(header_cell, "#4A4F55", diagnostics, header_location))
        body_fills.append(_cell_fill(body_cell, "#FFFFFF", diagnostics, body_location))
        columns.append(
            ColumnModel(
                name=name,
                heading=heading,
                alignment=body_style.alignment,
                format_mask=None,
                inferred_weight=weights[index - 1],
            )
        )

    def common_style(
        values: list[TextStyle],
        label: str,
        default: TextStyle,
        *,
        include_alignment: bool = True,
    ) -> TextStyle:
        if not values:
            return default
        signatures = {
            (
                v.font_family,
                v.font_size,
                v.font_weight,
                v.font_color,
                v.alignment if include_alignment else None,
            )
            for v in values
        }
        if len(signatures) > 1:
            diagnostics.warning(
                "DOCX-STYLE-005",
                f"{label} usa estilos mixtos; APEX requiere un estilo uniforme.",
                location="Tabla principal",
                strict_error=True,
            )
        return values[0]

    table_header = common_style(
        header_styles,
        "La fila de encabezado",
        TextStyle("HELVETICA", 9, "BOLD", "#FFFFFF", "CENTER"),
    )
    table_body = common_style(
        body_styles,
        "La fila prototipo",
        TextStyle("HELVETICA", 8.5, "NORMAL", "#25282B", "START"),
        include_alignment=False,
    )
    if len(set(header_fills)) > 1 or len(set(body_fills)) > 1:
        diagnostics.warning(
            "DOCX-STYLE-006",
            "Los fondos de la tabla no son uniformes.",
            location="Tabla principal",
            strict_error=True,
        )
    table_header = TextStyle(**{**table_header.to_dict(), "background_color": header_fills[0] if header_fills else "#4A4F55"})
    table_body = TextStyle(**{**table_body.to_dict(), "alignment": None, "background_color": body_fills[0] if body_fills else "#FFFFFF"})

    footer_paragraphs = section.footer.paragraphs
    nonempty_footer = [paragraph for paragraph in footer_paragraphs if paragraph.text.strip()]
    if len(nonempty_footer) > 1:
        diagnostics.error("DOCX-STRUCT-009", "El footer debe contener un solo párrafo con contenido.")
    footer = nonempty_footer[0] if nonempty_footer else footer_paragraphs[0]
    if _utf8_length(footer.text) > 4000:
        diagnostics.error(
            "DOCX-STRUCT-011",
            "El pie no puede superar 4000 bytes UTF-8 antes de sustituir campos.",
        )
    footer_fields = validate_text_placeholders(
        footer.text,
        allowed_kinds={"APP_USER", "GENERATED_AT", "FIELD"},
        diagnostics=diagnostics,
        location="Footer",
    )
    title_style = _paragraph_style(
        header,
        diagnostics,
        location="Párrafo de encabezado",
        default_size=15,
        default_color="#2F343A",
        default_bold=True,
        minimum_size=8,
        maximum_size=24,
        forced_alignment="CENTER",
    )
    footer_style = _paragraph_style(
        footer,
        diagnostics,
        location="Footer",
        default_size=8,
        default_color="#666666",
        default_bold=False,
        minimum_size=6,
        maximum_size=12,
    )
    border_width, border_color = _table_border(table, diagnostics)
    # Word puede guardar una página horizontal solo con sus dimensiones, sin
    # ``w:orient``; ambas formas deben producir la misma orientación.
    landscape = section.orientation == WD_ORIENT.LANDSCAPE or (
        section.page_width is not None
        and section.page_height is not None
        and section.page_width > section.page_height
    )
    orientation = "LANDSCAPE" if landscape else "PORTRAIT"
    fields = tuple(dict.fromkeys((*header_fields, *footer_fields)))
    _validate_a4(section, diagnostics)

    if diagnostics.has_errors:
        return None
    return TemplateModel(
        source_path=path,
        source_sha256=digest,
        header_template=header.text,
        footer_template=footer.text,
        fields=fields,
        columns=tuple(columns),
        template_orientation=orientation,
        title_style=title_style,
        table_header_style=table_header,
        table_body_style=table_body,
        border_width=border_width,
        border_color=border_color,
        footer_style=footer_style,
    )

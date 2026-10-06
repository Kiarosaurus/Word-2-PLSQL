#!/usr/bin/env python3
"""Genera la plantilla de demostración del modo layout: proyectos/estado_cuenta/estado_cuenta.docx.

Estado de cuenta ficticio de un alumno, con la apariencia típica de un
reporte de Oracle Reports y todo lo que el modo layout admite: tablas de
maquetación con bordes visibles, blancos o sin borde; cuatro tablas de datos
con cabecera repetida y totales; un recuadro de resumen; y encabezado y pie de
Word con número de página ({{PAGE}} DE {{PAGES}}).

Uso::

    python tools/generate_layout_template.py
"""

from __future__ import annotations

import importlib.util
from pathlib import Path
import shutil
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "proyectos" / "estado_cuenta" / "estado_cuenta.docx"
FONT = "Arial"
MONEY = "FM999G999G990D00"
BLACK, WHITE, GREY, BAND, LIGHT = "000000", "FFFFFF", "D9D9D9", "BFBFBF", "D9D9D9"

_spec = importlib.util.spec_from_file_location("sample_templates", ROOT / "tools" / "generate_sample_templates.py")
_sample = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = _sample
_spec.loader.exec_module(_sample)


# ---------------------------------------------------------------- utilidades
def run(paragraph, text: str, *, size: float = 7, bold: bool = False, italic: bool = False, color: str = BLACK):
    r = paragraph.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    fonts = r._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    return r


def tight(paragraph, align: str = "L", mark_size: float | None = None):
    paragraph.alignment = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER,
                           "R": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if mark_size is not None:
        p_pr = paragraph._p.get_or_add_pPr()
        r_pr = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(int(mark_size * 2)))
        r_pr.append(size)
        p_pr.append(r_pr)
    return paragraph


def text_cell(cell, lines, *, align: str = "L", size: float = 7, bold: bool = False, italic: bool = False):
    """lines: str con saltos \\n, o lista de fragmentos (texto, {opciones})."""
    paragraph = tight(cell.paragraphs[0], align, size)
    pieces = lines if isinstance(lines, list) else [(lines, {})]
    for text, options in pieces:
        parts = text.split("\n")
        for index, part in enumerate(parts):
            r = run(paragraph, part, size=options.get("size", size), bold=options.get("bold", bold),
                    italic=options.get("italic", italic))
            if index < len(parts) - 1:
                r.add_break()
    return paragraph


KEEP = object()
TC_AFTER_BORDERS = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
TC_AFTER_SHADING = TC_AFTER_BORDERS[1:]


def border(cell, *, top=KEEP, bottom=KEEP, left=KEEP, right=KEEP):
    """Cada lado: None (sin borde), (ancho_pt, color) o KEEP (no cambia).

    Reescribe tcBorders completo y en el orden del esquema OOXML.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    current = {}
    if old is not None:
        for child in old:
            current[child.tag] = child
        tc_pr.remove(old)
    node = OxmlElement("w:tcBorders")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if value is KEEP:
            if qn(f"w:{side}") in current:
                node.append(current[qn(f"w:{side}")])
            continue
        element = OxmlElement(f"w:{side}")
        if value is None:
            element.set(qn("w:val"), "nil")
        else:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(int(value[0] * 8)))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), value[1])
        node.append(element)
    tc_pr.insert_element_before(node, *TC_AFTER_BORDERS)


def box(cell, width: float = 0.5, color: str = BLACK):
    border(cell, top=(width, color), bottom=(width, color), left=(width, color), right=(width, color))


def fill(cell, color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, *TC_AFTER_SHADING)


def new_table(container, rows: int, widths_mm: list[float], *, center: bool = False, pad_mm: float = 0.8):
    if hasattr(container, "add_table") and container.__class__.__name__ == "Document":
        table = container.add_table(rows=rows, cols=len(widths_mm))
    else:
        table = container.add_table(rows=rows, cols=len(widths_mm), width=Mm(sum(widths_mm)))
    tbl_pr = table._tbl.tblPr
    # Orden del esquema: ... jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
    if center:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tbl_pr.insert_element_before(jc, "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                     "w:tblLayout", "w:tblCellMar", "w:tblLook")
    else:
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "0")
        indent.set(qn("w:type"), "dxa")
        tbl_pr.insert_element_before(indent, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 0), ("left", pad_mm), ("bottom", 0), ("right", pad_mm)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(int(value * 1440 / 25.4)))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.insert_element_before(margins, "w:tblLook")
    table.autofit = False
    for index, width in enumerate(widths_mm):
        table._tbl.tblGrid.gridCol_lst[index].set(qn("w:w"), str(int(width * 1440 / 25.4)))
        for row in table.rows:
            row.cells[index].width = Mm(width)
    for row in table.rows:
        for cell in row.cells:
            # La marca de párrafo de una celda vacía también ocupa altura en Word.
            tight(cell.paragraphs[0], mark_size=6)
    return table


def spacer(container, size: float = 4):
    tight(container.add_paragraph(), mark_size=size)


def data_table(document, *, title: str, title_span: int, headings: list[str], cells: list[str],
               widths: list[float], aligns: list[str], totals: dict[int, str]):
    """Tabla de datos: fila de banda, fila de títulos, fila que se repite y fila de totales."""
    table = new_table(document, 4 if totals else 3, widths)
    band = table.rows[0].cells[0].merge(table.rows[0].cells[title_span - 1])
    fill(band, BAND)
    text_cell(band, title, italic=True, size=7.5)
    for index, heading in enumerate(headings):
        cell = table.rows[1].cells[index]
        if heading:
            box(cell, 0.5)
            text_cell(cell, heading, align="C", italic=True)
        cell_body = table.rows[2].cells[index]
        if cells[index]:
            border(cell_body, bottom=(0.25, LIGHT))
            text_cell(cell_body, cells[index], align=aligns[index])
    if totals:
        for index, marker in totals.items():
            cell = table.rows[3].cells[index]
            fill(cell, GREY)
            box(cell, 0.5)
            text_cell(cell, marker, align="R")
    spacer(document, 5)


# ---------------------------------------------------------------- documento
def build() -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(15)
    section.top_margin = Mm(26)
    section.bottom_margin = Mm(50)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(44)

    # Encabezado: tabla invisible de tres celdas + título (se repite en cada página).
    header = section.header
    header.is_linked_to_previous = False
    header._element.remove(header.paragraphs[0]._p)
    top = new_table(header, 1, [60, 60, 60], pad_mm=0)
    text_cell(top.rows[0].cells[0], "{{FIELD:SISTEMA}}\n{{FIELD:MODULO}}\n{{FIELD:CODIGO}}", size=6)
    text_cell(top.rows[0].cells[2],
              "{{GENERATED_AT|DD/MM/YYYY}}\n{{GENERATED_AT|HH:MI:SS AM}}\n{{APP_USER}}", align="R", size=6)
    title = tight(header.add_paragraph(), "C")
    run(title, "{{REPORT_TITLE}}", size=10)
    tight(header.add_paragraph(), mark_size=2)

    # Pie: línea superior, usuario y "X DE Y" en casillas.
    footer = section.footer
    footer.is_linked_to_previous = False
    footer._element.remove(footer.paragraphs[0]._p)
    bottom = new_table(footer, 1, [78, 5, 7, 5, 66], pad_mm=0.5)
    cells = bottom.rows[0].cells
    for cell in cells:
        border(cell, top=(1.5, BLACK))
    text_cell(cells[0], "Generado por {{APP_USER}}")
    box(cells[1]); text_cell(cells[1], "{{PAGE}}", align="C")
    text_cell(cells[2], "DE", align="C")
    box(cells[3]); text_cell(cells[3], "{{PAGES}}", align="C")
    tight(footer.add_paragraph(), mark_size=2)

    # Cabecera del alumno: cuadrícula de 8 columnas (4 bloques etiqueta/valor).
    widths = [24, 21, 22, 23, 21, 20, 25, 22]
    grid = new_table(document, 6, widths)
    band = grid.rows[0].cells[0].merge(grid.rows[0].cells[7])
    fill(band, GREY)
    text_cell(band, [("Alumno   :  {{FIELD:ALUMNO.COD_ALUMNO}}          {{FIELD:ALUMNO.CARRERA}}          ",
                      {"size": 8.5}),
                     ("{{FIELD:ALUMNO.NOMBRE}}", {"size": 8.5, "italic": True})])
    blocks = [
        [("Código :", "{{FIELD:ALUMNO.COD_ALUMNO}}", "L"),
         ("Fecha Matrícula :", "{{FIELD:ALUMNO.FEC_MATRICULA|DD/MM/YYYY}}", "L"),
         ("Modalidad :", "{{FIELD:ALUMNO.MODALIDAD}}", "L"),
         ("Sede :", "{{FIELD:ALUMNO.SEDE}}", "L"), ("", "", "L")],
        [("Costo Programa :", "{{FIELD:ALUMNO.COSTO_PROGRAMA|%s}}" % MONEY, "R"),
         ("Saldo Anterior :", "{{FIELD:ALUMNO.SALDO_ANTERIOR|%s}}" % MONEY, "R"),
         ("Seguro :", "{{FIELD:ALUMNO.SEGURO|%s}}" % MONEY, "R"),
         ("Beca :", "{{FIELD:ALUMNO.BECA|%s}}" % MONEY, "R"),
         ("Importe Neto :", "{{FIELD:ALUMNO.IMPORTE_NETO|%s}}" % MONEY, "R")],
        [("Cuota Pensión :", "{{FIELD:ALUMNO.CUOTA_PENSION|%s}}" % MONEY, "R"),
         ("Cuota Servicios :", "{{FIELD:ALUMNO.CUOTA_SERVICIOS|%s}}" % MONEY, "R"),
         ("Cuota Mensual", "{{FIELD:ALUMNO.CUOTA_MENSUAL|%s}}" % MONEY, "R"),
         ("", "", "L"),
         ("Nro Cuotas :", "{{FIELD:ALUMNO.NRO_CUOTAS}}", "R")],
        [("Programa : S/.", "{{FIELD:ALUMNO.COSTO_PROGRAMA|%s}}" % MONEY, "R"),
         ("Servicios : S/.", "{{FIELD:ALUMNO.SERVICIOS_TOTAL|%s}}" % MONEY, "R"),
         ("Deuda Total : S/.", "{{FIELD:ALUMNO.DEUDA_TOTAL|%s}}" % MONEY, "R"),
         ("", "", "L"),
         ("Recargo Mensual %:", "{{FIELD:ALUMNO.RECARGO_MENSUAL|FM990D00}}", "R")],
    ]
    frame = (0.75, BLACK)
    for block_index, block in enumerate(blocks):
        label_col, value_col = block_index * 2, block_index * 2 + 1
        for row_offset, (label, value, align) in enumerate(block, start=1):
            label_cell = grid.rows[row_offset].cells[label_col]
            value_cell = grid.rows[row_offset].cells[value_col]
            text_cell(label_cell, label, size=7.5)
            text_cell(value_cell, value, size=7.5, align=align,
                      italic=(block_index == 0 and label == "Sede :"))
            last = row_offset == len(block)
            # Etiqueta/valor separados por un borde BLANCO (invisible); bloques
            # separados por un borde negro; sin borde horizontal entre filas.
            border(label_cell,
                   left=frame if block_index == 0 else (0.5, BLACK),
                   right=(0.5, WHITE),
                   bottom=frame if last else None)
            border(value_cell,
                   left=(0.5, WHITE),
                   right=frame if block_index == 3 else None,
                   bottom=frame if last else None)
    # Recuadro alrededor de cuota mensual y deuda total.
    border(grid.rows[3].cells[5], top=(0.5, BLACK), bottom=(0.5, BLACK), left=(0.5, BLACK), right=(0.5, BLACK))
    border(grid.rows[3].cells[7], top=(0.5, BLACK), bottom=(0.5, BLACK), left=(0.5, BLACK), right=frame)
    box(band, 0.75)
    spacer(document, 6)

    data_table(
        document, title="PAGOS REGISTRADOS EN CAJA", title_span=5,
        headings=["Año", "Mes", "Caja", "N° Cuota", "Pensión", "Servicios", "Descuento", "Mora",
                  "TOTAL", "Boleta", "Importe"],
        cells=["{{COLUMN:CAJA.ANIO}}", "{{COLUMN:CAJA.MES}}", "{{COLUMN:CAJA.CAJA}}",
               "{{COLUMN:CAJA.NRO_CUOTA}}", "{{COLUMN:CAJA.PENSION|%s}}" % MONEY,
               "{{COLUMN:CAJA.SERVICIOS|%s}}" % MONEY, "{{COLUMN:CAJA.DESCUENTO|%s}}" % MONEY,
               "{{COLUMN:CAJA.MORA|%s}}" % MONEY, "{{COLUMN:CAJA.TOTAL|%s}}" % MONEY,
               "{{COLUMN:CAJA.BOLETA}}", "{{COLUMN:CAJA.IMPORTE|%s}}" % MONEY],
        widths=[7, 11.4, 14, 13.9, 18, 16, 16, 14, 20, 24, 20],
        aligns=["L", "C", "C", "C", "R", "R", "R", "R", "R", "C", "R"],
        totals={4: "{{SUM:CAJA.PENSION|%s}}" % MONEY, 5: "{{SUM:CAJA.SERVICIOS|%s}}" % MONEY,
                7: "{{SUM:CAJA.MORA|%s}}" % MONEY, 8: "{{SUM:CAJA.TOTAL|%s}}" % MONEY},
    )
    data_table(
        document, title="CUOTAS POR PAGAR", title_span=5,
        headings=["Año", "Mes", "Caja", "N° Cuota", "Pensión", "Servicios", "", "TOTAL"],
        cells=["{{COLUMN:CUOTAS.ANIO}}", "{{COLUMN:CUOTAS.MES}}", "{{COLUMN:CUOTAS.CAJA}}",
               "{{COLUMN:CUOTAS.NRO_CUOTA}}", "{{COLUMN:CUOTAS.PENSION|%s}}" % MONEY,
               "{{COLUMN:CUOTAS.SERVICIOS|%s}}" % MONEY, "", "{{COLUMN:CUOTAS.TOTAL|%s}}" % MONEY],
        widths=[7, 11.4, 14, 13.9, 18, 16, 30, 20],
        aligns=["L", "C", "C", "R", "R", "R", "L", "R"],
        totals={4: "{{SUM:CUOTAS.PENSION|%s}}" % MONEY, 5: "{{SUM:CUOTAS.SERVICIOS|%s}}" % MONEY,
                7: "{{SUM:CUOTAS.TOTAL|%s}}" % MONEY},
    )
    data_table(
        document, title="PAGOS POR TRANSFERENCIA BANCARIA", title_span=5,
        headings=["N°", "Fecha", "N° Operación", "N° Cuota", "Pensión", "Servicios", "Descuento", "Mora", "TOTAL"],
        cells=["{{COLUMN:TRANSFERENCIAS.NRO}}", "{{COLUMN:TRANSFERENCIAS.FECHA|DD/MM/YYYY}}",
               "{{COLUMN:TRANSFERENCIAS.NRO_OPERACION}}", "{{COLUMN:TRANSFERENCIAS.NRO_CUOTA}}",
               "{{COLUMN:TRANSFERENCIAS.PENSION|%s}}" % MONEY, "{{COLUMN:TRANSFERENCIAS.SERVICIOS|%s}}" % MONEY,
               "{{COLUMN:TRANSFERENCIAS.DESCUENTO|%s}}" % MONEY, "{{COLUMN:TRANSFERENCIAS.MORA|%s}}" % MONEY,
               "{{COLUMN:TRANSFERENCIAS.TOTAL|%s}}" % MONEY],
        widths=[7, 18.4, 21, 13.9, 18, 16, 16, 14, 20],
        aligns=["C", "C", "C", "C", "R", "R", "R", "R", "R"],
        totals={4: "{{SUM:TRANSFERENCIAS.PENSION|%s}}" % MONEY, 5: "{{SUM:TRANSFERENCIAS.SERVICIOS|%s}}" % MONEY,
                6: "{{SUM:TRANSFERENCIAS.DESCUENTO|%s}}" % MONEY, 7: "{{SUM:TRANSFERENCIAS.MORA|%s}}" % MONEY,
                8: "{{SUM:TRANSFERENCIAS.TOTAL|%s}}" % MONEY},
    )
    data_table(
        document, title="MESES CON PAGO ATRASADO", title_span=5,
        headings=["Año", "Mes", "", "N°", "", "", "", "Importe"],
        cells=["{{COLUMN:ATRASOS.ANIO}}", "{{COLUMN:ATRASOS.MES}}", "", "{{COLUMN:ATRASOS.NRO}}", "", "", "",
               "{{COLUMN:ATRASOS.IMPORTE|%s}}" % MONEY],
        widths=[7, 11.4, 14, 13.9, 18, 16, 30, 20],
        aligns=["L", "C", "L", "R", "L", "L", "L", "R"],
        totals={7: "{{SUM:ATRASOS.IMPORTE|%s}}" % MONEY},
    )

    # Recuadro de resumen: tabla centrada con borde exterior y sin bordes internos.
    rows = [
        ("Costo total del programa", "+", "COSTO_TOTAL"),
        ("Pagos registrados en caja", "-", "PAGOS_CAJA"),
        ("Becas y descuentos", "-", "BECAS"),
        ("Pagos por transferencia bancaria", "-", "PAGOS_TRANSFERENCIA"),
        ("Otros abonos", "-", "OTROS_ABONOS"),
        ("Devoluciones", "+", "DEVOLUCIONES"),
    ]
    summary = new_table(document, len(rows) + 2, [85, 12, 28], center=True, pad_mm=1.5)
    head = summary.rows[0].cells[0].merge(summary.rows[0].cells[2])
    text_cell(head, "RESUMEN DE LA CUENTA DEL ALUMNO", align="C", size=8, bold=True)
    for index, (label, sign, column) in enumerate(rows, start=1):
        cells = summary.rows[index].cells
        text_cell(cells[0], label, size=7.5)
        text_cell(cells[1], sign, align="C", size=7.5)
        text_cell(cells[2], "{{FIELD:RESUMEN.%s|%s}}" % (column, MONEY), align="R", size=7.5)
    last = summary.rows[len(rows) + 1].cells
    text_cell(last[0], "Saldo Pendiente", align="R", size=8, italic=True)
    text_cell(last[1], "S/", align="C", size=8, italic=True)
    text_cell(last[2], "{{FIELD:RESUMEN.SALDO|%s}}" % MONEY, align="R", size=8, bold=True)
    outer = (0.75, BLACK)
    count = len(summary.rows)
    for r_index, row in enumerate(summary.rows):
        seen = set()
        for c_index, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            keep_bottom = r_index == len(rows) and c_index == 2       # raya antes del saldo
            border(cell,
                   top=outer if r_index == 0 else None,
                   bottom=outer if r_index == count - 1 else ((0.75, BLACK) if keep_bottom else None),
                   left=outer if c_index == 0 else None,
                   right=outer if c_index == 2 or r_index == 0 else None)
    tight(document.add_paragraph(), mark_size=2)

    properties = document.core_properties
    properties.title = "Estado de cuenta del alumno (ejemplo ficticio del modo layout)"
    properties.author = "Área de Informática"
    properties.created = _sample.FIXED_TIMESTAMP
    properties.modified = _sample.FIXED_TIMESTAMP
    return document


def main() -> int:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as directory:
        raw = Path(directory) / "raw.docx"
        clean = Path(directory) / "clean.docx"
        build().save(raw)
        _sample._normalize_docx_archive(raw, clean)       # quita customXml de python-docx
        shutil.copyfile(clean, OUTPUT)
        shutil.copyfile(clean, ROOT / "templates" / "estado_cuenta_layout.docx")   # referencia del modo layout
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

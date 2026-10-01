#!/usr/bin/env python3
"""Genera las plantillas DOCX de referencia del compilador.

Las plantillas resultantes implementan únicamente el lenguaje restringido de
``docs/CONTRACT.md``. No contienen imágenes, campos de Word, macros, formas ni
contenido ejecutable.

Uso::

    python tools/generate_sample_templates.py
    python tools/generate_sample_templates.py --output-dir examples/templates

El script fija los metadatos y las marcas de tiempo del contenedor ZIP para que
la salida sea reproducible entre ejecuciones con la misma versión de
``python-docx``.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "examples" / "templates"
FIXED_TIMESTAMP = datetime(2026, 1, 1, 0, 0, 0, tzinfo=UTC)
ZIP_TIMESTAMP = (2026, 1, 1, 0, 0, 0)
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


@dataclass(frozen=True, slots=True)
class ColumnSpec:
    """Definición visual de una columna en la plantilla de referencia."""

    heading: str
    alias: str
    width_cm: float
    body_alignment: WD_ALIGN_PARAGRAPH


@dataclass(frozen=True, slots=True)
class TemplateSpec:
    """Configuración completa de una plantilla reproducible."""

    filename: str
    orientation: str
    margin_cm: float
    body_font_size_pt: float
    columns: tuple[ColumnSpec, ...]


PORTRAIT_SPEC = TemplateSpec(
    filename="entidades_portrait.docx",
    orientation="PORTRAIT",
    margin_cm=1.8,
    body_font_size_pt=8.5,
    columns=(
        ColumnSpec("DNI", "VDNI", 3.2, WD_ALIGN_PARAGRAPH.CENTER),
        ColumnSpec("Nombre completo", "VNOM", 8.2, WD_ALIGN_PARAGRAPH.LEFT),
        ColumnSpec("Departamento", "DEPARTAMENTO", 5.6, WD_ALIGN_PARAGRAPH.LEFT),
    ),
)


LANDSCAPE_SPEC = TemplateSpec(
    filename="entidades_landscape.docx",
    orientation="LANDSCAPE",
    margin_cm=1.4,
    body_font_size_pt=8.0,
    columns=(
        ColumnSpec("DNI", "VDNI", 2.4, WD_ALIGN_PARAGRAPH.CENTER),
        ColumnSpec("Nombre completo", "VNOM", 4.9, WD_ALIGN_PARAGRAPH.LEFT),
        ColumnSpec("Dirección actual", "VDIREC_ACTUAL", 6.0, WD_ALIGN_PARAGRAPH.LEFT),
        ColumnSpec("Departamento", "DEPARTAMENTO", 4.8, WD_ALIGN_PARAGRAPH.LEFT),
        ColumnSpec("Distrito", "DISTRITO", 3.2, WD_ALIGN_PARAGRAPH.LEFT),
        ColumnSpec("Teléfono", "VNRO_TLF1", 4.6, WD_ALIGN_PARAGRAPH.CENTER),
    ),
)


def _set_run_style(
    run: object,
    *,
    size_pt: float,
    bold: bool,
    color: str,
) -> None:
    """Aplica un estilo que el compilador puede traducir sin pérdida."""

    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    # Word puede ignorar ``font.name`` si no se especifican ambos atributos.
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")


def _set_cell_fill(cell: object, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell: object, *, top: int, start: int, bottom: int, end: int) -> None:
    """Establece padding en twips sin introducir elementos no admitidos."""

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: object, *, color: str = "BFC3C7", size: int = 4) -> None:
    """Define un borde uniforme de 0.5 pt en toda la tabla."""

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_grid_widths(table: object, widths_cm: Iterable[float]) -> None:
    """Sincroniza ``tblGrid`` y las celdas para inferencia estable de pesos."""

    widths = tuple(widths_cm)
    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    for index, width_cm in enumerate(widths):
        width = Cm(width_cm)
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(width.twips))
        table.columns[index].width = width
        for row in table.rows:
            row.cells[index].width = width


def _set_repeat_table_header(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def _remove_title_border(document: Document, paragraph: object) -> None:
    """Elimina la línea azul heredada del estilo ``Title`` de Word.

    La plantilla conserva el estilo semántico ``Title``, pero neutraliza el
    borde inferior tanto en la definición del estilo como en el párrafo. Esto
    evita diferencias entre Word y LibreOffice y refleja la capacidad real de
    APEX_DATA_EXPORT, que no representa esa línea decorativa.
    """

    title_style = document.styles["Title"]
    style_p_pr = title_style.element.get_or_add_pPr()
    style_borders = style_p_pr.find(qn("w:pBdr"))
    if style_borders is not None:
        style_p_pr.remove(style_borders)

    paragraph_p_pr = paragraph._p.get_or_add_pPr()
    paragraph_borders = paragraph_p_pr.find(qn("w:pBdr"))
    if paragraph_borders is not None:
        paragraph_p_pr.remove(paragraph_borders)


def _configure_document(document: Document, spec: TemplateSpec) -> None:
    """Construye el único layout permitido por el contrato 1.0."""

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    if spec.orientation == "LANDSCAPE":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    else:
        section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(spec.margin_cm)
    section.right_margin = Cm(spec.margin_cm)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    # El header real permanece vacío. APEX_DATA_EXPORT usa el único párrafo
    # previo a la tabla como encabezado textual del reporte.
    title = document.add_paragraph(style="Title")
    _remove_title_border(document, title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("{{REPORT_TITLE}}")
    _set_run_style(title_run, size_pt=15, bold=True, color="2F343A")

    table = document.add_table(rows=2, cols=len(spec.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _set_grid_widths(table, (column.width_cm for column in spec.columns))
    _set_repeat_table_header(table.rows[0])

    for index, column in enumerate(spec.columns):
        header_cell = table.cell(0, index)
        body_cell = table.cell(1, index)
        for cell in (header_cell, body_cell):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=90, start=110, bottom=90, end=110)

        header_cell.text = ""
        header_paragraph = header_cell.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_paragraph.paragraph_format.space_before = Pt(0)
        header_paragraph.paragraph_format.space_after = Pt(0)
        header_run = header_paragraph.add_run(column.heading)
        _set_run_style(header_run, size_pt=9, bold=True, color="FFFFFF")
        _set_cell_fill(header_cell, "4A4F55")

        body_cell.text = ""
        body_paragraph = body_cell.paragraphs[0]
        body_paragraph.alignment = column.body_alignment
        body_paragraph.paragraph_format.space_before = Pt(0)
        body_paragraph.paragraph_format.space_after = Pt(0)
        body_run = body_paragraph.add_run(f"{{{{COLUMN:{column.alias}}}}}")
        # El tamaño es uniforme en toda la fila porque APEX_DATA_EXPORT aplica
        # un único estilo al cuerpo. En landscape se usa 8 pt y se reservan
        # anchos mayores para los aliases largos; no se altera ningún alias.
        _set_run_style(
            body_run,
            size_pt=spec.body_font_size_pt,
            bold=False,
            color="25282B",
        )
        _set_cell_fill(body_cell, "FFFFFF")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer_run = footer.add_run("Usuario: {{APP_USER}} | Fecha: {{GENERATED_AT}}")
    _set_run_style(footer_run, size_pt=8, bold=False, color="666666")

    properties = document.core_properties
    properties.title = "Plantilla de reporte corporativo"
    properties.subject = "Plantilla restringida para Oracle APEX"
    properties.author = "Área de Informática"
    properties.last_modified_by = "Área de Informática"
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP


def _normalize_docx_archive(source: Path, destination: Path) -> None:
    """Reescribe el ZIP de Word con orden, timestamps y partes permitidas.

    La plantilla base distribuida con ``python-docx`` incluye una parte
    ``customXml`` de bibliografía aunque el documento no la utilice. El
    compilador aplica una política de rechazo estricto a ``customXml``; por
    eso el generador elimina esa parte y sus referencias antes de publicar la
    plantilla de ejemplo.
    """

    def sanitized_part(name: str, data: bytes) -> bytes:
        if name == "word/_rels/document.xml.rels":
            root = ET.fromstring(data)
            for relationship in list(root):
                if relationship.attrib.get("Type", "").endswith("/customXml"):
                    root.remove(relationship)
            ET.register_namespace("", REL_NS)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        if name == "[Content_Types].xml":
            root = ET.fromstring(data)
            for child in list(root):
                if child.attrib.get("PartName", "").startswith("/customXml/"):
                    root.remove(child)
            ET.register_namespace("", CT_NS)
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return data

    with zipfile.ZipFile(source, "r") as input_zip, zipfile.ZipFile(
        destination,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as output_zip:
        for name in sorted(input_zip.namelist()):
            if name.casefold().startswith("customxml/"):
                continue
            source_info = input_zip.getinfo(name)
            info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = source_info.external_attr
            info.flag_bits = source_info.flag_bits & ~0x08
            output_zip.writestr(info, sanitized_part(name, input_zip.read(name)))


def generate_template(spec: TemplateSpec, output_dir: Path) -> Path:
    """Genera una plantilla y la publica atómicamente en ``output_dir``."""

    output_dir.mkdir(parents=True, exist_ok=True)
    destination = output_dir / spec.filename
    with tempfile.TemporaryDirectory(prefix="apex-report-template-") as temp_dir:
        raw_path = Path(temp_dir) / spec.filename
        normalized_path = Path(temp_dir) / f"normalized-{spec.filename}"
        document = Document()
        _configure_document(document, spec)
        document.save(raw_path)
        _normalize_docx_archive(raw_path, normalized_path)
        shutil.copyfile(normalized_path, destination)
    return destination


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera las dos plantillas DOCX de referencia del compilador.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directorio de salida (predeterminado: {DEFAULT_OUTPUT_DIR}).",
    )
    return parser.parse_args()


def main() -> int:
    arguments = parse_args()
    for spec in (PORTRAIT_SPEC, LANDSCAPE_SPEC):
        output = generate_template(spec, arguments.output_dir.resolve())
        print(f"{output}  sha256={_sha256(output)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

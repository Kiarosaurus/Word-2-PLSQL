"""Convierte documentación Markdown en documentos Word con el estilo corporativo.

Las fuentes viven en ``tools/doc_sources``. Además del Markdown habitual
(títulos, listas, tablas y bloques de código) se admite un bloque de diagrama
de flujo que se dibuja con tablas nativas de Word, sin imágenes::

    ```flujo
    # Leyenda opcional del diagrama
    Actor | Paso | Detalle opcional
    ? Actor | ¿Pregunta de decisión? | Sí: continúa · No: vuelve al paso 1
    ```

Cada línea es una caja; entre cajas consecutivas se dibuja una flecha. Una
línea que empieza con ``?`` se dibuja como decisión.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables"
NAVY = "283848"
SLATE = "4A5563"
LIGHT = "E8EDF2"
PALE = "F5F7F9"
WHITE = "FFFFFF"
AMBER = "B7791F"
AMBER_PALE = "FDF3E1"
TEXT = RGBColor(37, 43, 49)
SOURCES = ROOT / "tools" / "doc_sources"


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_keep(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_inline(paragraph, text: str, *, code: bool = False) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<![*\w])\*[^*\s][^*]*\*(?![*\w]))")
    for part in pattern.split(text):
        if not part:
            continue
        italic = part.startswith("*") and not part.startswith("**") and part.endswith("*") and len(part) > 2
        run = paragraph.add_run(part[1:-1] if part.startswith("`") and part.endswith("`") else
                                part[2:-2] if part.startswith("**") and part.endswith("**") else
                                part[1:-1] if italic else part)
        if italic:
            run.italic = True
        if code or (part.startswith("`") and part.endswith("`")):
            run.font.name = "Consolas"
            run.font.size = Pt(8.2)
            run.font.color.rgb = RGBColor(35, 54, 72)
        if part.startswith("**") and part.endswith("**"):
            run.bold = True


def configure_document(document: Document, title: str, subtitle: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 24, NAVY), ("Heading 1", 17, NAVY),
                              ("Heading 2", 13, SLATE), ("Heading 3", 10.5, SLATE)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(subtitle)
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(SLATE)
    sub.paragraph_format.space_after = Pt(18)

    border = document.add_table(rows=1, cols=1)
    border.autofit = True
    shade(border.cell(0, 0), NAVY)
    border.cell(0, 0).height = Cm(0.11)
    document.add_paragraph()

    header = section.header.paragraphs[0]
    header.text = "APEX WORD REPORT COMPILER  ·  DOCUMENTACIÓN 1.0"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(SLATE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Uso interno · Oracle APEX 24.2")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(SLATE)


def add_code_block(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE)
    set_cell_margins(cell, 100, 140, 100, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        add_inline(p, line, code=True)


def set_cell_borders(cell, color: str | None, *, size: int = 8, style: str = "single") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        if color is None:
            node.set(qn("w:val"), "nil")
        else:
            node.set(qn("w:val"), style)
            node.set(qn("w:sz"), str(size))
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), color)


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def add_flow(document: Document, lines: list[str]) -> None:
    """Dibuja un diagrama de flujo vertical: columna de actor + columna de cajas."""

    caption = None
    steps: list[tuple[bool, str, str, str]] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            caption = line.lstrip("#").strip()
            continue
        decision = line.startswith("?")
        parts = [part.strip() for part in line.lstrip("?").split("|")]
        if len(parts) == 1:
            parts = ["", parts[0]]
        actor, title, detail = (parts + [""])[:3]
        steps.append((decision, actor, title, detail))
    if not steps:
        return

    table = document.add_table(rows=len(steps) * 2 - 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders_none(table)
    widths = (Cm(3.3), Cm(11.2))
    for row in table.rows:
        prevent_row_split(row)
        for index, width in enumerate(widths):
            row.cells[index].width = width

    for position, (decision, actor, title, detail) in enumerate(steps):
        row = table.rows[position * 2]
        actor_cell, box = row.cells
        actor_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(actor_cell, 40, 60, 40, 140)
        p = actor_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(actor.upper())
        run.font.name = "Aptos"
        run.font.size = Pt(7)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(AMBER if decision else SLATE)

        box.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(box, AMBER_PALE if decision else LIGHT)
        set_cell_borders(box, AMBER if decision else NAVY, size=10 if decision else 8,
                         style="dashed" if decision else "single")
        set_cell_margins(box, 70, 140, 70, 140)
        p = box.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, ("◆ " if decision else "") + title)
        for item in p.runs:
            item.bold = True
            item.font.size = Pt(8.8)
            if item.font.name != "Consolas":
                item.font.color.rgb = RGBColor.from_string(AMBER if decision else NAVY)
        if detail:
            d = box.add_paragraph()
            d.alignment = WD_ALIGN_PARAGRAPH.CENTER
            d.paragraph_format.space_after = Pt(0)
            add_inline(d, detail)
            for item in d.runs:
                item.font.size = Pt(7.8)

        if position < len(steps) - 1:
            arrow_row = table.rows[position * 2 + 1]
            arrow = arrow_row.cells[1].paragraphs[0]
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow.paragraph_format.space_after = Pt(0)
            arrow.paragraph_format.line_spacing = 0.9
            run = arrow.add_run("▼")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(SLATE)
            if decision:
                # La flecha que sale de una decisión es la rama afirmativa;
                # la negativa se describe en el detalle de la caja.
                label = arrow.add_run("  Sí")
                label.bold = True
                label.font.size = Pt(7.5)
                label.font.color.rgb = RGBColor.from_string(AMBER)

    # El diagrama se mantiene completo en una página.
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    if caption:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(3)
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(SLATE)
    else:
        document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_header(table.rows[0])
    for row_index, source_row in enumerate(rows):
        prevent_row_split(table.rows[row_index])
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            shade(cell, NAVY if row_index == 0 else WHITE if row_index % 2 else PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            value = source_row[col_index] if col_index < len(source_row) else ""
            add_inline(p, value)
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    if len(rows) <= 15:
        # Tablas breves: se mantienen en una sola página para no dejar filas
        # huérfanas; las largas se parten repitiendo el encabezado.
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    result: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            result.append(cells)
        index += 1
    return result, index


LIST_ITEM = re.compile(r"^(?P<indent>\s*)(?P<marker>[-*]|\d+\.)\s+(?P<text>.+)$")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
AUTOLINK = re.compile(r"<(https?://[^>]+)>")


def normalize_links(text: str) -> str:
    text = LINK.sub(lambda match: f"{match.group(1)} ({match.group(2)})", text)
    return AUTOLINK.sub(lambda match: match.group(1), text)


def is_block_start(stripped: str) -> bool:
    return (
        not stripped
        or stripped.startswith(("```", "|", "#"))
        or LIST_ITEM.match(stripped) is not None
    )


def add_list_item(document: Document, marker: str, text: str, level: int) -> None:
    indent = 0.65 + 0.6 * level
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(3)
    # Viñeta o número explícito: evita que Word continúe automáticamente una
    # lista anterior situada varias páginas antes.
    p.add_run(("• " if marker in {"-", "*"} else f"{marker} "))
    add_inline(p, text)


def markdown_to_docx(source: Path, destination: Path, title: str, subtitle: str) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document, title, subtitle)
    index = 0
    skipped_title = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            is_flow = stripped[3:].strip().lower() == "flujo"
            code_lines: list[str] = []
            fence_indent = len(raw) - len(raw.lstrip())
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                line = lines[index]
                code_lines.append(line[fence_indent:] if line[:fence_indent].isspace() else line)
                index += 1
            if is_flow:
                add_flow(document, code_lines)
            else:
                add_code_block(document, code_lines)
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, [[normalize_links(cell) for cell in row] for row in rows])
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if not skipped_title and level == 1:
                skipped_title = True
                index += 1
                continue
            paragraph = document.add_paragraph(style=f"Heading {min(max(level - 1, 1), 3)}")
            add_inline(paragraph, text)
            set_keep(paragraph)
            index += 1
            continue
        item = LIST_ITEM.match(raw)
        if item:
            parts = [item.group("text").strip()]
            index += 1
            while index < len(lines):
                following = lines[index]
                if is_block_start(following.strip()):
                    break
                parts.append(following.strip())
                index += 1
            level = min(len(item.group("indent")) // 2, 3)
            add_list_item(document, item.group("marker"), normalize_links(" ".join(parts)), level)
            continue

        # Párrafo: une las líneas ajustadas a mano; dos espacios finales
        # marcan un salto de línea explícito.
        p = document.add_paragraph()
        first = True
        while index < len(lines):
            line = lines[index]
            if not first and is_block_start(line.strip()):
                break
            if not first:
                p.add_run(" ")
            add_inline(p, normalize_links(line.strip()))
            first = False
            index += 1
            if line.endswith("  "):
                p.add_run().add_break()
                if index < len(lines) and not is_block_start(lines[index].strip()):
                    continue
                break

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = title
    document.core_properties.subject = subtitle
    document.core_properties.author = "Área de Informática"
    document.core_properties.keywords = "Oracle APEX, reportes, Word, APEX_DATA_EXPORT"
    document.save(destination)


# (origen en tools/doc_sources, destino Word, título, subtítulo)
DOCUMENTS = (
    ("MANUAL.md", "deliverables/Manual_de_uso.docx", "Manual de uso",
     "Compilador local Word → reportes Oracle APEX 24.2"),
    ("INFORME_ALCANCE.md", "deliverables/Informe_de_alcance_y_limitaciones.docx",
     "Informe de alcance y limitaciones",
     "Arquitectura, seguridad y fronteras del producto · Versión 1.0"),
    ("CONTRACT.md", "docs/CONTRACT.docx", "Contrato funcional y técnico",
     "Compilador Word restringido para reportes Oracle APEX 24.2"),
    ("TEMPLATE_QA.md", "docs/TEMPLATE_QA.docx", "Control de calidad de las plantillas",
     "Plantillas de referencia del contrato 1.0"),
    ("RELEASE_REVIEW.md", "docs/RELEASE_REVIEW.docx", "Revisión de liberación",
     "APEX Word Report Compiler 1.0.0"),
    ("README.md", "README.docx", "Compilador local Word → reportes Oracle APEX",
     "Guía de inicio · Versión 1.0"),
    ("SQL_README.md", "sql/README.docx", "Instalación de PKG_CORPORATE_REPORTS",
     "Package común para Oracle APEX 24.2"),
    ("EXAMPLES_README.md", "examples/README.docx", "Ejemplo ejecutable ENTIDADES",
     "Proyecto completo de referencia"),
)


def main(argv: list[str] | None = None) -> int:
    """Convierte Markdown a Word.

    Sin argumentos convierte la lista ``DOCUMENTS`` desde ``tools/doc_sources``.
    Con argumentos: ``origen.md destino.docx "Título" "Subtítulo"``.
    """

    import sys

    args = sys.argv[1:] if argv is None else argv
    if args:
        if len(args) != 4:
            print("Uso: build_documentation.py origen.md destino.docx \"Título\" \"Subtítulo\"")
            return 2
        markdown_to_docx(Path(args[0]), Path(args[1]), args[2], args[3])
        print(args[1])
        return 0
    for source, destination, title, subtitle in DOCUMENTS:
        if (SOURCES / source).is_file():
            markdown_to_docx(SOURCES / source, ROOT / destination, title, subtitle)
            print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
